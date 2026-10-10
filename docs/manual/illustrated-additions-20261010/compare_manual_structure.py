"""Read-only integrity gate for the illustrated manual update.

Uses integration_notes.json table_items as exact BOM allowances. It does not
render, edit either DOCX, update fields, or start Word. All indices are zero-based.
"""
from __future__ import annotations

import argparse
import copy
import difflib
import hashlib
import json
import re
import sys
from collections import Counter
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn

EMU_INCH = 914400
EMU_PT = 12700
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
      'wp': 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing',
      'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}


def sha(data):
    return hashlib.sha256(data).hexdigest()


def media(path):
    with ZipFile(path) as z:
        return {name: sha(z.read(name)) for name in z.namelist()
                if name.startswith('word/media/') and not name.endswith('/')}


def tables(doc):
    return [[[c.text for c in row.cells] for row in table.rows] for table in doc.tables]


def steps(doc):
    return [{'text': p.text, 'style': p.style.name}
            for p in doc.paragraphs if re.match(r'^步骤\s+\d{3}(?:\s|$)', p.text)]


def apply_table_allowances(original, notes, errors):
    result = copy.deepcopy(original)
    cell_changes = []
    inserted_rows = []
    insertions = []
    for item in notes.get('table_items', []):
        ti, ri = item['table_index'], item['row_index']
        try:
            actual = original[ti][ri]
        except IndexError:
            errors.append(f"{item['id']}: source table/row index is invalid")
            continue
        if actual != item['old_exact_paragraphs']:
            errors.append(f"{item['id']}: source row differs from exact authorized old text")
            continue
        replacement = item['recommended_adjacent_note_or_correction']
        if isinstance(replacement, list):
            if len(replacement) != len(actual):
                errors.append(f"{item['id']}: replacement column count changed")
                continue
            result[ti][ri] = replacement
            for ci, (old, new) in enumerate(zip(actual, replacement)):
                if old != new:
                    cell_changes.append({'id': item['id'], 'table': ti, 'row': ri,
                                         'column': ci, 'old': old, 'new': new})
        elif isinstance(replacement, dict) and replacement.get('operation') == 'insert_row_after':
            insertions.append((ti, ri, replacement['new_row'], item['id']))
        else:
            errors.append(f"{item['id']}: unrecognized table allowance")
    for ti, ri, row, identifier in sorted(insertions, reverse=True):
        result[ti].insert(ri + 1, row)
        inserted_rows.append({'id': identifier, 'table': ti, 'after_source_row': ri, 'values': row})
    return result, cell_changes, inserted_rows


def picture_dimensions(doc, original_hashes):
    """Map each drawing to its actual section; report geometric and caption risks."""
    layouts = []
    for section in doc.sections:
        layouts.append({'width': int(section.page_width - section.left_margin - section.right_margin),
                        'height': int(section.page_height - section.top_margin - section.bottom_margin)})
    images = []
    section_no = 0
    for child in doc.element.body:
        layout = layouts[min(section_no, len(layouts) - 1)]
        context = ''.join(child.itertext())[:160]
        for drawing in child.iter(qn('w:drawing')):
            extent = next(drawing.iter(qn('wp:extent')), None)
            blip = next(drawing.iter(qn('a:blip')), None)
            if extent is None or blip is None:
                continue
            rel = blip.get(qn('r:embed'))
            if not rel or rel not in doc.part.related_parts:
                continue
            blob = doc.part.related_parts[rel].blob
            digest = sha(blob)
            width, height = int(extent.get('cx')), int(extent.get('cy'))
            hard_risks = []
            warnings = []
            if width > layout['width'] + EMU_PT:
                hard_risks.append('Image width exceeds section text width')
            if height > layout['height'] + EMU_PT:
                hard_risks.append('Image height exceeds section text height')
            if height > layout['height'] - 32 * EMU_PT:
                warnings.append('Less than 32 pt remains for caption and paragraph spacing; visually inspect this page')
            desc = next(drawing.iter(qn('wp:docPr')), None)
            images.append({'sequence': len(images) + 1, 'section': section_no + 1,
                           'new_media': digest not in original_hashes, 'sha256': digest,
                           'width_in': round(width / EMU_INCH, 3),
                           'height_in': round(height / EMU_INCH, 3),
                           'usable_width_in': round(layout['width'] / EMU_INCH, 3),
                           'usable_height_in': round(layout['height'] / EMU_INCH, 3),
                           'alt': desc.get('descr', '') if desc is not None else '',
                           'nearby_text': context, 'hard_risks': hard_risks, 'warnings': warnings})
        if child.tag == qn('w:p') and child.find('./w:pPr/w:sectPr', NS) is not None:
            section_no += 1
    return images


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--original', required=True, type=Path)
    ap.add_argument('--updated', required=True, type=Path)
    ap.add_argument('--notes', type=Path, default=Path(__file__).with_name('integration_notes.json'))
    ap.add_argument('--report', type=Path, default=Path(__file__).with_name('structure_audit.json'))
    ap.add_argument('--expected-original-images', type=int, default=38)
    ap.add_argument('--expected-added-images', type=int, default=9)
    ap.add_argument('--expected-added-tables', type=int, default=7,
                    help='Appendix tables after the preserved original 43 tables')
    ap.add_argument('--baseline-selfcheck', action='store_true',
                    help='Compare an unchanged baseline to itself; skip BOM allowances and require zero additions')
    args = ap.parse_args()
    old = Document(args.original)
    new = Document(args.updated)
    notes = json.loads(args.notes.read_text(encoding='utf-8'))
    errors = []
    warnings = []
    old_steps, new_steps = steps(old), steps(new)
    expected_numbers = list(range(1, 331))
    for label, data in [('original', old_steps), ('updated', new_steps)]:
        numbers = [int(re.search(r'\d{3}', item['text']).group()) for item in data]
        if numbers != expected_numbers:
            errors.append(f'{label}: step numbers are not exactly 001 through 330')
        if any(item['style'] != 'Heading 3' for item in data):
            errors.append(f'{label}: a step title is not Heading 3')
    if old_steps != new_steps:
        errors.append('Original 330 step titles/styles were changed')

    old_media, new_media = media(args.original), media(args.updated)
    old_counts, new_counts = Counter(old_media.values()), Counter(new_media.values())
    missing_media = old_counts - new_counts
    added_media = new_counts - old_counts
    expected_added = 0 if args.baseline_selfcheck else args.expected_added_images
    if len(old.inline_shapes) != args.expected_original_images or len(old_media) != args.expected_original_images:
        errors.append('Original embedded/inline image count differs from expected 38')
    if len(new.inline_shapes) != args.expected_original_images + expected_added:
        errors.append('Updated inline image count differs from expected total')
    if len(new_media) != args.expected_original_images + expected_added:
        errors.append('Updated embedded media count differs from expected total')
    if missing_media:
        errors.append('An original media SHA256 (including multiplicity) is missing or changed')
    if sum(added_media.values()) != expected_added:
        errors.append('Embedded media SHA256 additions do not equal expected count')

    old_tables, new_tables = tables(old), tables(new)
    expected_added_tables = 0 if args.baseline_selfcheck else args.expected_added_tables
    if args.baseline_selfcheck:
        expected_tables, cell_changes, row_insertions = old_tables, [], []
    else:
        expected_tables, cell_changes, row_insertions = apply_table_allowances(old_tables, notes, errors)
        if len(cell_changes) != 3 or len(row_insertions) != 1:
            errors.append('Authorized BOM delta must be exactly three changed cells and one inserted row')
    table_diffs = []
    if len(old_tables) != 43:
        errors.append('Original baseline does not contain exactly 43 tables')
    if len(new_tables) != len(old_tables) + expected_added_tables:
        errors.append('Updated table count is not original 43 plus the expected seven appendix tables')
    for ti in range(len(expected_tables)):
        expected = expected_tables[ti]
        actual = new_tables[ti] if ti < len(new_tables) else None
        if expected != actual:
            table_diffs.append({'table': ti, 'expected': expected, 'actual': actual})
    if table_diffs:
        errors.append('Tables differ beyond the exact authorized BOM cell edits and inserted row')
    added_table_records = []
    for ti in range(len(old_tables), len(new_tables)):
        data = new_tables[ti]
        added_table_records.append({'table': ti, 'rows': len(data),
                                    'columns_per_row': [len(row) for row in data],
                                    'header': data[0] if data else [], 'content': data})
        if len(data) < 2 or not any(cell.strip() for row in data for cell in row):
            errors.append(f'Appendix table {ti} has no meaningful header and data rows')

    dimensions = picture_dimensions(new, set(old_media.values()))
    for item in dimensions:
        if item['new_media'] and item['hard_risks']:
            errors.append(f"New image {item['sequence']}: " + '; '.join(item['hard_risks']))
        if item['new_media'] and item['warnings']:
            warnings.append({'image': item['sequence'], 'warnings': item['warnings']})

    def body_paragraphs(doc):
        return [p.text for p in doc.paragraphs if not p.style.name.lower().startswith('toc')]
    old_paras, new_paras = body_paragraphs(old), body_paragraphs(new)
    paragraph_delta = []
    for tag, a, b, c, d in difflib.SequenceMatcher(a=old_paras, b=new_paras, autojunk=False).get_opcodes():
        if tag != 'equal':
            paragraph_delta.append({'operation': tag, 'original_range': [a, b], 'updated_range': [c, d],
                                    'old': old_paras[a:b], 'new': new_paras[c:d]})
    report = {'status': 'PASS' if not errors else 'FAIL',
              'original': str(args.original.resolve()), 'updated': str(args.updated.resolve()),
              'original_sha256': sha(args.original.read_bytes()), 'updated_sha256': sha(args.updated.read_bytes()),
              'step_titles_identical': old_steps == new_steps, 'step_count': len(new_steps),
              'original_images': len(old.inline_shapes), 'updated_images': len(new.inline_shapes),
              'original_embedded_media': len(old_media), 'updated_embedded_media': len(new_media),
              'missing_original_media_sha256': dict(missing_media), 'added_media_sha256': dict(added_media),
              'original_tables': len(old_tables), 'updated_tables': len(new_tables),
              'authorized_cell_changes': cell_changes, 'authorized_row_insertions': row_insertions,
              'unexpected_table_diffs': table_diffs,
              'added_appendix_tables': added_table_records,
              'added_table_scope': 'Seven expected: two SG90 tables, one chapter index, two board component tables, two board wire tables.',
              'new_image_dimensions': [i for i in dimensions if i['new_media']],
              'paragraph_delta_informational': paragraph_delta,
              'paragraph_gate_note': 'Step headings are exact; body paragraph changes are reported for review because the 19 integration notes and cover count edits are authorized.',
              'errors': errors, 'warnings': warnings,
              'visual_qa_still_required': True}
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({key: report[key] for key in ['status','step_count','updated_images','updated_tables','errors','warnings']}, ensure_ascii=False))
    return 1 if errors else 0


if __name__ == '__main__':
    if hasattr(sys.stdout, 'reconfigure'):
        sys.stdout.reconfigure(encoding='utf-8')
    raise SystemExit(main())
