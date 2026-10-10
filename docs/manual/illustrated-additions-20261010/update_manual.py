"""Append verified linkage and perfboard references to the retained illustrated manual.

Inputs are deliberately explicit. The baseline is the retained 139-page enhanced
manual, not the older 125-page original. Run only once against that baseline.
"""
from pathlib import Path
from copy import deepcopy
import argparse, csv, json, re, hashlib
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
NAME = 'STM32与RFID智能移动垃圾桶_逐步组装与调试实施手册_图解增强版.docx'
ap = argparse.ArgumentParser()
ap.add_argument('--input', type=Path, default=HERE/'before.docx')
ap.add_argument('--output', type=Path, default=HERE/'staging'/NAME)
ap.add_argument('--perfboard', type=Path, default=ROOT/'SmartBin_洞洞板接线图_20261010')
ap.add_argument('--notes', type=Path, default=ROOT/'.illustrated-manual-20261010/external_wiring_notes.json')
args = ap.parse_args()
doc = Document(args.input)
baseline_steps = [p.text for p in doc.paragraphs if re.match(r'步骤 \d{3}', p.text)]
notes = json.loads((HERE/'integration_notes.json').read_text(encoding='utf-8'))
layout = json.loads((args.perfboard/'layout.json').read_text(encoding='utf-8'))
audit = {'baseline_sha256': hashlib.sha256(args.input.read_bytes()).hexdigest(), 'edits': []}

def format_run(r, size=None, bold=None):
    r.font.name = '微软雅黑'
    r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '微软雅黑')
    if size is not None: r.font.size = Pt(size)
    if bold is not None: r.bold = bold
    r.font.color.rgb = RGBColor(0, 0, 0)

def replace_text(p, text):
    saved = deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    p.clear()
    r = p.add_run(text)
    if saved is not None: r._r.insert(0, saved)

def insert_after(anchor, text):
    e = OxmlElement('w:p'); anchor._p.addnext(e)
    p = Paragraph(e, anchor._parent); p.style = doc.styles['Normal']
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    format_run(p.add_run(text), 10)
    return p

def exact(text):
    out = [p for p in doc.paragraphs if p.text == text]
    if len(out) != 1: raise ValueError(f'Anchor count {len(out)}: {text[:100]}')
    return out[0]

for item in notes['paragraph_items']:
    if item['id'] == 'INT01':
        replacements = [
            '保留原版330个步骤与33项功能，现含39幅操作图解及1幅整机外观概念图；原有图1至图7继续保留。',
            '图解01—30位于原操作章节；图解31—35为SG90桶盖机构详图，图解36—39为PWR与IO洞洞板正反面图，集中见第60—61章。原有图1至图7、330个步骤和接线表保留。先看元件面识别孔位，再按已镜像的焊接面及逐根焊线表施工，最后补齐外置模块线束。'
        ]
    else: replacements = item['recommended_adjacent_note_or_correction']
    if item['id'] == 'INT17':
        replacements = replacements.replace('的两根原短线', '的全部原连线').replace('同步修改新IO参考图的相关线记录', '按第61.8节重接旁路并修改新IO参考图的相关线记录')
    if item['id'] == 'INT14': replacements = replacements.replace('SERVO5V','SERVO_SIG_5V')
    if item['operation'].startswith('replace'):
        if isinstance(replacements, str): replacements = [replacements]
        for old, new in zip(item['old_exact_paragraphs'], replacements):
            replace_text(exact(old), new)
            audit['edits'].append({'id': item['id'], 'old': old, 'new': new})
    else:
        text = replacements
        if item['id'] == 'INT02': text += ' 详见第61章图解36—39及焊线表。'
        if item['id'] == 'INT14': text += ' 机械连接图见第60章图解31—35。'
        insert_after(exact(item['old_exact_paragraphs'][0]), text)
        audit['edits'].append({'id': item['id'], 'insert': text})

for item in notes['table_items']:
    table = doc.tables[item['table_index']]
    found = [r for r in table.rows if [c.text for c in r.cells] == item['old_exact_paragraphs']]
    if len(found) != 1: raise ValueError(f'BOM anchor {item["id"]}: {len(found)}')
    row = found[0]; spec = item['recommended_adjacent_note_or_correction']
    if isinstance(spec, dict):
        new_xml = deepcopy(row._tr); row._tr.addnext(new_xml)
        from docx.table import _Row
        new_row = _Row(new_xml, table)
        for c, value in zip(new_row.cells, spec['new_row']):
            replace_text(c.paragraphs[0], value)
            for p in c.paragraphs[1:]: p._p.getparent().remove(p._p)
    else:
        for c, value in zip(row.cells, spec): replace_text(c.paragraphs[0], value)
    audit['edits'].append({'id': item['id'], 'table_change': spec})

anchor = exact('图解 09  先量两孔中心距，再装连接端；球头应可摆动，维护拆端头连接螺钉即可。')
insert_after(anchor, '合盖、中间、全开三个状态，以及侧臂穿盖叠层和舵盘球头固定的放大图，见第60章图解31—35；该章沿用本节20—50—20—50mm机构尺寸。')

def para(text, size=10.5, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = keep
    p.paragraph_format.widow_control = True
    format_run(p.add_run(text), size)
    return p

def heading(text, level=2, newpage=False):
    p = doc.add_paragraph(text, style=f'Heading {level}')
    if newpage: p.paragraph_format.page_break_before = True
    p.paragraph_format.keep_with_next = True
    for r in p.runs: r.font.color.rgb = RGBColor(0, 0, 0)
    return p

def figure(path, caption, alt, width):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.keep_together = True
    run = p.add_run(); run.add_picture(str(path), width=Cm(width))
    for pr in run._r.xpath('.//wp:docPr'): pr.set('descr', alt)
    c = para(caption, 9.5)
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.keep_together = True
    c.paragraph_format.space_after = Pt(5)
    audit.setdefault('figures', []).append({'caption': caption, 'file': path.name, 'width_cm': width})

def table(headers, rows, widths, size=9.5, center=()):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, widths): col.width = Cm(width)
    for cell, width in zip(t.rows[0].cells, widths): cell.width = Cm(width)
    for cell, value in zip(t.rows[0].cells, headers): cell.text = str(value)
    for vals in rows:
        row = t.add_row()
        for cell, width, value in zip(row.cells, widths, vals):
            cell.width = Cm(width); cell.text = str(value)
    pr = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e = OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
    pr.append(borders)
    for i, row in enumerate(t.rows):
        trPr = row._tr.get_or_add_trPr(); trPr.append(OxmlElement('w:cantSplit'))
        if i == 0: trPr.append(OxmlElement('w:tblHeader'))
        for j, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for side, value in [('top',65),('bottom',65),('left',80),('right',80)]:
                m = OxmlElement('w:'+side); m.set(qn('w:w'),str(value)); m.set(qn('w:type'),'dxa'); margins.append(m)
            cp.append(margins)
            if i == 0:
                shade = OxmlElement('w:shd'); shade.set(qn('w:fill'),'E9EFF5'); cp.append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.05
                p.paragraph_format.keep_with_next = i == 0
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 or j in center else WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs: format_run(r, size, i == 0)
    return t

heading('60  SG90舵机与桶盖连接详图', 1, True)
para('本章把第15章的连杆装配和第29章的舵机接线放大说明。先核对机械关系与三根线的去向，再按原步骤056—066及149—151装配、校准。本章各图是尺寸关系示意，不能按打印大小直接划孔。')
heading('60.1  先认清固定件和运动件')
para('SG90壳体固定在桶外支架上；桶盖绕真实铰链轴转动。原配舵盘带动20mm舵臂，经孔心距50mm的双球头拉杆，推动与桶盖刚性连接的20mm侧臂。舵机轴和盖铰链轴同高、平行，轴间距50mm，形成不交叉的平行四边形。')
para('闭盖时桶盖定义为0°，两臂相对水平为20°；正常开盖70°时，两臂相对水平为90°。这些是机构的几何角度，不是可直接写进程序的SG90角度指令。舵机先卸臂以50Hz、1500μs空载置中，再按原步骤小幅校准端点，禁止硬顶机械止挡。')
table(['核对项','本方案数值或要求'],[
    ('两固定轴间距','50mm；以真实轴心测量，两轴同高且平行'),
    ('舵机臂和盖侧臂','各20mm；均从对应轴心量到球头中心'),
    ('双球头拉杆','两球头中心距50mm；长度不是金属杆外露长度'),
    ('桶盖与活动总成','桶盖160×150×2mm；桶盖、随盖侧臂和紧固件总质量实称≤45g'),
    ('正常运动范围','盖0°→70°；先断执行电源并拆一端拉杆，才可维护撑盖'),
], [4.5,13.4], center=(0,))
heading('60.2  三根舵机线分别接到哪里')
table(['SG90线','连接位置','核对方法'],[
    ('红色电源线','PWR的SERVO支路5V','来自5V降压分支；不接8.4V电池或3.3V'),
    ('棕色或黑色地线','直接回PWR的SERVO支路GND','承载舵机电流，不能经IO小信号地线回流'),
    ('橙 黄 白信号线','IO的JACT中SERVO_SIG_5V，孔X9','信号名中的5V指PWM逻辑；来自AHCT第3脚Y1'),
    ('核心PA0','AHCT第2脚A1','经缓冲后才去SG90；两板和核心必须共地'),
], [3.1,7.0,7.8])
para('延长线按实物逐根测通断并贴标签；颜色仅作辅助。先不装拉杆验证空载，再安装舵盘及连杆。舵机运动和轮电机行走仍按原手册分时互锁。')
sg = json.loads((HERE/'sg90/manifest.json').read_text(encoding='utf-8'))
titles = ['合盖状态与尺寸基准','中间状态与干涉检查','正常全开与端点校准','盖侧臂与桶盖固定','舵盘与球头叠装']
for i, (item, title) in enumerate(zip(sg, titles), 31):
    heading(f'60.{i-28}  {title}', 2, True)
    figure(HERE/'sg90'/item['file'], f'图解 {i:02d}  '+item['caption'], item['alt'], item['suggested_width_cm'])
    para(item['explanation'])
    if i == 34:
        para('先用纸样核对舵机和铰链的真实位置，再确定侧臂固定孔。盖侧臂不能只靠一个可转动螺钉连接桶盖；两点固定和内侧加强片共同防转、分散受力。每次增加活动紧固件后重新称重。')
    if i == 35:
        para('装完先断电手动检查行程，再空载慢速试开。球头不能卡死或碰盖边；防松螺母应起防松作用，不能通过松螺母让夹层摇晃来冒充球头转动。连续20次无卡滞后，再继续原手册的负载和整机验收。')

heading('61  洞洞板正反面与逐根接线表', 1, True)
para('本章是两块7×9cm洞洞板的参考施工布局R1，适用于STM32F103C8T6核心板和已焊好的MCP23017扩展模块。图、孔位表和142条板内连线记录相互对应；尚未用你的实物模块焊接、通电验证，不能据此省略原手册的测量与验收。这些是手工洞洞板参考图，不是可直接打板的Gerber。')
heading('61.1  读图前先确认板型和方向')
intro = [
    '先确认2.54mm独立焊盘，在实物上找到连续24列×30行有效孔区，避开安装孔、螺钉及边缘缺孔。从元件面定义A—X和1—30，在A1旁画三角。有效孔中心跨度58.42×73.66mm；70×90mm外形不保证所有商品板孔区相同。',
    '元件面从左到右A—X；保持第1行朝上，绕竖轴翻到焊接面后，从左到右为X—A。图解37和39已经镜像，不能再镜像一次。两面同名孔始终是同一个物理孔。',
    '方框焊盘表示排针第1脚；端子正负和三极管E/B/C已标注。独立焊盘本身不相通，靠规定导线建立网络。整页图用于总览，小脚号看不清时放大SVG或查本章文字孔位表，不凭线条颜色猜接。',
    '先空插端子、IC座、电阻和电容，核对脚距与壳体间隙。U1是DIP14座，排距7.62mm、缺口朝第1行；电阻建议本体不大于约7×3mm，图示脚距10.16mm；电解脚距5.08mm、直径不大于10mm，正负按实际标识。',
    '142条记录由135条绝缘飞线和两根汇流铜线的7段记录组成。PWR-W001—W003是一根连续正铜线V3→V9，W004—W007是一根连续负铜线X3→X11；用约0.8mm镀锡实心铜线，不能切成七根细跳线。',
    '其余飞线仅在指定两端剥皮焊接；折线是建议绕线路径，交叉不相连。焊一根、断电测一根、勾一根。功率线按线表20AWG，其他按22/26AWG；大电流不依赖细焊锡桥。',
    'PWR板的5V_BUS与IO的5V_LOGIC为同一个5V节点；3V3_EXT与3V3_CORE不可直接连接，GND共地。电机和舵机直接回PWR，不经IO排针。核心、MCP成品模块、驱动和稳压模块均在板外固定。',
    'PWR放ADC分压与照明电阻；AO3400A转接板用核实后的G/S/D线束外接。IO图中的Q1/Q2/Q3对应原步骤中的QES1/QES2/QES3。图中接口次序是板侧自定义，不能按此猜模块原排针。',
]
for s in intro:
    p = para(s, 10)
    p.paragraph_format.line_spacing = 1.10
    p.paragraph_format.space_after = Pt(5)
heading('61.2  怎样从原步骤查到本章')
table(['原章节或步骤','本章使用位置'],[
    ('步骤094与图解03','先读本章方向规则，再看图解36—39；原图解03只讲通用焊接方法'),
    ('步骤096—114','用PWR/IO孔位表定位电源、核心线束、I²C和MCP接口'),
    ('步骤127 134 149 169—183','同时看条件上拉、QES位号、舵机支路、外置AO和PWR ADC说明'),
    ('板内线全部完成后','继续61.9外置模块接线；142条板内线不包括全部模块线束'),
], [5.4,12.5])

perf = [('PWR_front','PWR电源板元件面','从元件侧观察，A列在左，元件脚号以孔位表为准。'),
        ('PWR_back','PWR电源板焊接面','已经左右镜像，X列在左；两根裸汇流铜线检查后覆盖绝缘罩。'),
        ('IO_front','IO信号板元件面','AHCT缺口朝第1行，外置MCP成品模块不装入DIP28座。'),
        ('IO_back','IO信号板焊接面','已经左右镜像；跨线绝缘，焊点只在指定的起点和终点。')]
for number, (base, title, caution) in enumerate(perf, 36):
    heading(f'61.3  {title}' if number == 36 else title, 2, True)
    path = HERE/'perfboard-print'/f'{base}_print.png'
    figure(path, f'图解 {number:02d}  {title}。{caution}', title+'，24列30行完整孔区，保持第1行朝上。', 17.5)

for board_id, subsection in [('PWR','61.4'),('IO','61.5')]:
    heading(f'{subsection}  {board_id}元件孔位与接口', 2, board_id == 'PWR')
    para('引脚表示电气脚名，孔位始终采用元件面定义。网络名必须一致；NC表示留空。接口连接外部模块时按模块丝印逐根核对。装配说明中的条件元件必须根据实物上拉和电平选择，不能全部无条件焊满。', 10)
    parts = layout['boards'][board_id]['components']
    rows = []
    for c in parts:
        pins = '\n'.join(f'{p["pin"]} → {p["hole"]} → {p["net"]}' for p in c['pins'])
        rows.append((c['ref'],c['value'],pins,c.get('notes','')))
    table(['位号','规格或用途','引脚 → 孔位 → 网络','装配说明'],rows,[1.4,3.4,6.0,7.1],9.0,center=(0,))

for board_id, subsection in [('PWR','61.6'),('IO','61.7')]:
    heading(f'{subsection}  {board_id}逐根焊线检查表',2,True)
    para('断电逐行测量起终点导通，并确认未误接相邻网络后勾选。孔号不随翻面变化。表中两处“0.8mm汇流铜线”网络按连续整线制作，其他全部使用绝缘导线。',10)
    rows = list(csv.DictReader((args.perfboard/f'{board_id}_逐根焊线表.csv').open(encoding='utf-8-sig', newline='')))
    entries = [(r['焊线编号'].replace(board_id+'-',''),r['网络'],r['起点孔']+' / '+r['起点元件脚'],r['终点孔']+' / '+r['终点元件脚'],r['线规'],'□') for r in rows]
    table(['编号','网络','起点孔与元件脚','终点孔与元件脚','线规','检查'],entries,[1.6,3.2,4.4,4.4,3.3,1.0],9.0,center=(0,5))

heading('61.8  条件装配与外部线束边界',2)
para('以下条件决定是否改变默认图，须在焊接前写在纸质表旁。修改默认接线时同步划去被拆除的线记录；不能把原线和选配线叠加。')
para('蜂鸣器需要5V高电平时：默认AHCT通道3的9脚接GND、10脚接5V_LOGIC，且这两个孔还承担网络中转。必须拆除它们原来的全部相关线，再补回其他节点间的电源或地连接；不能仅在原线之上叠接。已确认兼容3.3V控制的三线蜂鸣模块直接用PB9，不做本选配改线。')
para('选配改线清单：先断电并拔下U1，拆掉IO-W005、W010、W011、W038、W041共5条默认线。另加D10→G12、D10→G6两根地旁路线，加A21→G5一根5V旁路线，保留原W039和W040。这样其他地和AHCT电源网络仍然连通。')
para('然后另接G8→D10（U1第10脚/OE3接地），G9→核心PB9（第9脚A3），G10→蜂鸣IN（第8脚Y3）。这6条新增连接均用26AWG绝缘线；断电检查G9已不再接地、G8已不再接5V，G4/G5仍接5V、G6/G12仍接地。确认无电源短路后插回U1，再按步骤173检查不响偏置与上电默认状态。默认142条线只对应未启用本分支的布局。')
para('编码器CLK/DT上拉只选一组：移除模块原上拉后装R22/R23拉3V3_CORE；或保留已确认拉3V3_EXT的模块上拉，不装R22/R23并验证电平。不能经两组上拉把两路3.3V稳压互接。I²C、RESET、片选和DHT上拉同样要计入模块已有元件。')
para('三路普通按键与蜂鸣偏置、每个5V数字输入的条件适配、5V DHT的双向适配、MCP与DFPlayer就近去耦、全部外置模块连接，不包含在142条板内焊线中。只完成两块洞洞板不能视为完成全机接线。')

heading('61.9  外置模块接线与分阶段上电',2)
external = json.loads(args.notes.read_text(encoding='utf-8'))['sections']
for idx, section in enumerate(external,1):
    heading(f'{idx:02d}  '+section['title'],3)
    for i, text in enumerate(section['lines'],1):
        # Keep conditional pull-up handling consistent with the corrected core text.
        if section['title'].startswith('编码器') and i == 1:
            text = '编码器确认支持3.3V后，VCC→3V3_EXT，CLK→PB3，DT→PB4，SW→GPB4，GND共地。CLK/DT只保留一组已核对电源轨的上拉：模块自带上拉拉3V3_EXT时不装IO的R22/R23；移除模块上拉后才用R22/R23拉3V3_CORE。SW约10kΩ拉3V3_EXT或使用已验证的MCP内部上拉。'
        para(f'{i}. {text}',10)

heading('61.10  文件和本次修订记录',2)
para('配套可放大SVG、高清PNG、孔位CSV、焊线CSV及layout.json随项目仓库保留在hardware/current-low-cost/perfboard/。本Word内已收录四张板图、全部65个位号的孔位与装配说明，以及51条PWR、91条IO连线记录，离线装配无需另找图表。')
para('2026年10月10日更新：加入5幅SG90详图与4幅洞洞板图；保留原330个步骤。局部补正AHCT第三通道拆线条件、编码器上拉选择、ADC位于PWR、AO转接板外置和参考孔区定义；增加0.8mm汇流铜线采购项。图纸审查通过不代表实物或程序已验收，仍执行原逐步检查。')

assert [p.text for p in doc.paragraphs if re.match(r'步骤 \d{3}',p.text)] == baseline_steps
assert len(baseline_steps)==330
assert len(doc.inline_shapes)==47, len(doc.inline_shapes)
args.output.parent.mkdir(parents=True,exist_ok=True)
doc.save(args.output)
audit.update({'output':str(args.output),'pictures':len(doc.inline_shapes),'tables':len(doc.tables),'step_count':330,'board_parts':{k:len(v['components']) for k,v in layout['boards'].items()}})
(HERE/'authoring_audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in audit.items() if k not in ['edits','figures']},ensure_ascii=False))
