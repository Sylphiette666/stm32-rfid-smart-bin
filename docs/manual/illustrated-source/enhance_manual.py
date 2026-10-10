"""Insert instructional figures without rebuilding or renumbering the source manual."""
from pathlib import Path
from copy import deepcopy
import argparse, json, re, hashlib
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

HERE=Path(__file__).parent
parser=argparse.ArgumentParser()
parser.add_argument('--input',type=Path,default=HERE.parent/'STM32与RFID智能移动垃圾桶_逐步组装与调试实施手册.docx')
parser.add_argument('--output',type=Path,default=HERE.parent/'STM32与RFID智能移动垃圾桶_逐步组装与调试实施手册_图解增强版.docx')
local_concept=HERE.parent/'stm32-rfid-smart-bin/docs/appearance/smart-bin-concept-v1.png'
repo_concept=HERE.parent.parent/'appearance/smart-bin-concept-v1.png'
parser.add_argument('--concept',type=Path,default=local_concept if local_concept.exists() else repo_concept)
args=parser.parse_args()
original_hash=hashlib.sha256(args.input.read_bytes()).hexdigest()
doc=Document(args.input)
original_paras=[p.text for p in doc.paragraphs]
original_steps=[p.text for p in doc.paragraphs if re.match(r'步骤 \d{3}',p.text)]
original_tables=[[[c.text for c in row.cells] for row in t.rows] for t in doc.tables]
original_picture_count=len(doc.inline_shapes)

def new_para(anchor,text='',after=False):
    e=OxmlElement('w:p')
    if after:anchor._p.addnext(e)
    else:anchor._p.addprevious(e)
    p=Paragraph(e,anchor._parent);p.style=doc.styles['Normal']
    if text:p.add_run(text)
    return p

def figure(anchor,file,caption,alt,after=False):
    p=new_para(anchor,after=after)
    p.paragraph_format.keep_with_next=True
    p.paragraph_format.keep_together=True
    p.paragraph_format.space_before=Pt(7)
    p.paragraph_format.space_after=Pt(3)
    p.paragraph_format.line_spacing=1
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=p.add_run();run.add_picture(str(file),width=Cm(17.79))
    for pr in run._r.xpath('.//wp:docPr'):pr.set('descr',alt)
    c=new_para(p,caption,after=True)
    c.paragraph_format.keep_together=True
    c.paragraph_format.keep_with_next=False
    c.paragraph_format.space_after=Pt(9)
    c.paragraph_format.line_spacing=1.08
    c.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in c.runs:r.font.size=Pt(9.5);r.font.color.rgb=RGBColor.from_string('444444')
    return c

items=[]
for category in ['general','mechanical','electrical','software']:
    p=HERE/f'{category}_figures.json'
    if not p.exists():raise RuntimeError(f'Missing required figure manifest {p}')
    for item in json.loads(p.read_text(encoding='utf-8')):
        exact=[(i,p) for i,p in enumerate(doc.paragraphs) if p.text==item['anchor']]
        if len(exact)!=1:raise RuntimeError(f'Anchor count {len(exact)} for {item["id"]}: {item["anchor"]}')
        index,anchor=exact[0]
        path=HERE/item['file']
        if not path.is_file():raise FileNotFoundError(path)
        items.append((index,category,anchor,item,path))
items.sort(key=lambda v:(v[0],v[3]['id']))
audit_items=[]
for n,(index,category,anchor,item,path) in enumerate(items,1):
    caption=f'图解 {n:02d}  '+item['caption']
    figure(anchor,path,caption,item['alt'],item.get('position','before')=='after')
    audit_items.append(dict(number=n,category=category,**item))

# Local additions to the front matter; original titles, chapters and operation text stay in place.
subtitle=next(p for p in doc.paragraphs if p.style.name=='Subtitle')
subtitle.text='逐步组装与调试实施手册  图解增强版'
cover_date=next(p for p in doc.paragraphs if p.text=='装配版本 2026年10月10日')
note=new_para(cover_date,f'保留原版330个步骤与33项功能，新增{len(items)}幅操作图解及1幅整机外观概念图。',after=True)
for r in note.runs:r.font.size=Pt(10)
toc=next(p for p in doc.paragraphs if p.text=='目录')
intro=new_para(toc,'整机外观与图解读法')
intro.paragraph_format.page_break_before=True
intro.paragraph_format.keep_with_next=True
for r in intro.runs:r.bold=True;r.font.size=Pt(18);r.font.color.rgb=RGBColor(0,0,0)
context=new_para(toc,'先看整机由哪些部分组成，再按正文步骤装配。下面的外观概念图用于理解两层底盘、桶体和开盖机构的关系；板件尺寸、孔位与电气连接以正文标注为准。')
context.paragraph_format.keep_with_next=True
figure(toc,args.concept,'外观概念图  合盖与开盖状态。AI生成示意，不是实物照片或加工模板。','整机合盖与开盖的AI外观概念图，显示木质底盘、白色桶壳、读卡区和外置舵机。')
for text in [
    '看图顺序：先读图题，沿数字或箭头找操作顺序，再读图下说明，最后完成紧邻的步骤和“□ 检查”。图中绿色表示通过现象，红色可表示电源正极或需要处理的问题，必须结合文字判断。',
    '机械图按毫米标注，剖面和局部会放大；长度以数字为准，不按打印图比例切割。接线图中的模块框按信号名排列，不能当作实物排针顺序。',
    '原有图1至图7保留；新增图解按01起连续编号，分布在基础操作、机械、电气和调试章节。新增图片是操作辅助，原有接线表、实测要求和330个步骤继续有效。'
]:
    p=new_para(toc,text)
    p.paragraph_format.space_after=Pt(8)

# Tighten only reference-entry spacing so the final two links do not occupy
# an otherwise empty page. Retain the original text and type sizes.
in_references=False
for p in doc.paragraphs:
    if p.text.startswith('E-S1 STMicroelectronics'):
        in_references=True
    if in_references:
        p.paragraph_format.space_before=Pt(0)
        p.paragraph_format.space_after=Pt(3 if p.text.startswith('http') else 0)

doc.core_properties.title='STM32与RFID智能移动垃圾桶逐步组装与调试实施手册图解增强版'
doc.core_properties.subject='原330步组装调试手册的机械接线测试图解增强版本'
doc.core_properties.comments='保留原手册，添加操作示意图与图下注释。图示不代表实物验收已完成。'
args.output.parent.mkdir(parents=True,exist_ok=True)
doc.save(args.output)

# Preservation checks deliberately compare operation text and complete table contents.
out=Document(args.output)
assert [p.text for p in out.paragraphs if re.match(r'步骤 \d{3}',p.text)]==original_steps
assert len(original_steps)==330
assert [[[c.text for c in row.cells] for row in t.rows] for t in out.tables]==original_tables
assert len(out.inline_shapes)==original_picture_count+len(items)+1
assert hashlib.sha256(args.input.read_bytes()).hexdigest()==original_hash
orig_content=[p.text for p in Document(args.input).paragraphs if not p.style.name.startswith('toc') and p.style.name!='Subtitle']
after_content=[p.text for p in out.paragraphs]
cursor=0
for text in orig_content:
    while cursor<len(after_content) and after_content[cursor]!=text:cursor+=1
    assert cursor<len(after_content),f'Original paragraph missing: {text[:120]}'
    cursor+=1
report=dict(input=str(args.input),input_sha256=original_hash,output=str(args.output),steps=330,functions=33,original_figures=original_picture_count,added_operation_figures=len(items),added_concept_figures=1,total_figures=len(out.inline_shapes),tables=len(out.tables),preserved_original_text=True,figures=audit_items)
(HERE/'enhancement_audit.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in report.items() if k!='figures'},ensure_ascii=False))
