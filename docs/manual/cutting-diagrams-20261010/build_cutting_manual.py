from pathlib import Path
from copy import deepcopy
import json, re, hashlib, zipfile, argparse
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

HERE=Path(__file__).resolve().parent
ROOT=HERE.parent
NAME='STM32与RFID智能移动垃圾桶_逐步组装与调试实施手册_图解增强版.docx'
ap=argparse.ArgumentParser()
ap.add_argument('--input',type=Path,default=HERE/'before.docx')
ap.add_argument('--output',type=Path,default=HERE/'staging'/NAME)
args=ap.parse_args()
doc=Document(args.input)
audit={'baseline_sha256':hashlib.sha256(args.input.read_bytes()).hexdigest(),'figures':[]}
assert audit['baseline_sha256']=='251cd1a64bbd1a466630d4e966b3866fa1ca21587e8ba4a98dcf5b7da0dbdf59'
steps_before=[p.text for p in doc.paragraphs if re.match(r'步骤 \d{3}',p.text)]
tables_before=[[[c.text for c in r.cells] for r in t.rows] for t in doc.tables]
assert len(steps_before)==330 and len(doc.inline_shapes)==47 and len(doc.tables)==50

# Reuse the formatting helpers from the previous retained-manual edit, not its builder.
exec(compile((HERE/'retained_formatting.py').read_text(encoding='utf-8'),'<retained-formatting>','exec'))

entries=[]
for group in ('wood','pvc'):
    items=json.loads((HERE/group/'figure_manifest.json').read_text(encoding='utf-8'))
    if isinstance(items,dict): items=items.get('figures',items.get('items'))
    for item in items:
        item['group']=group
        entries.append(item)
entries.sort(key=lambda x:int(x['number']))
assert [int(x['number']) for x in entries]==list(range(40,45))

for p in doc.paragraphs:
    if p.text.startswith('保留原版330个步骤与33项功能，现含39幅'):
        replace_text(p,'保留原版330个步骤与33项功能，现含44幅操作图解及1幅整机外观概念图；原有图1至图7继续保留。')
    elif p.text.startswith('图解01—30位于原操作章节；图解31—35'):
        replace_text(p,'图解01—30位于原操作章节；图解31—35为SG90连接详图，图解36—39为洞洞板正反面图，图解40—44为木板与PVC板裁剪图，集中见第60—62章。原有图1至图7、330个步骤和接线表保留。先核对材料厚度与成品尺寸，再划线裁剪；未知安装孔均按实物试装后描孔。')

heading('62  木板与PVC板裁剪图解',1,True)
para('本章对应原步骤015—020、036—039、050—059。先完成纸样和材料核对，再切外形，最后按实物确定安装孔。各图均为尺寸示意，单位为mm；按标注数值量取，不能按图片或打印后的大小直接描边。')
para('木件使用两张300×300×5胶合板；固定桶壳与附件使用3mm PVC发泡板（雪弗板）；活动盖使用另购的2mm PVC发泡板。1mm PP板是内胆材料，不替代本章PVC板。现有3mm PVC共两片，按排料图使用并保留剩料。')
table(['材料与零件','成品或毛坯尺寸','数量','加工要求'],[
('5mm木板 P01底板','280×220','1','外形切准，四角去锐边约C2'),
('5mm木板 P02上平台','240×190','1','安装时相对P01偏移(20,15)'),
('5mm木板 舵机立架','60×190','1','排料时可旋转；装机时190为高'),
('5mm木板 调节托片','约80×40','1','外形毛坯；安装/检修孔先试装'),
('3mm PVC P03左右侧壁','160×180','2','180为竖直高度'),
('3mm PVC P04前后壁','144×180','2','夹在侧壁之间，不切成150宽'),
('3mm PVC P05壳底','154×144','1','嵌在四壁内，底面与壁板底缘齐平'),
('3mm PVC P07操作板','190×120毛坯','1','按实际模块开窗、避让开盖空间'),
('3mm PVC P08超声横梁','150×35毛坯','1','可拆装；传感器与端耳按实物定孔'),
('3mm PVC P10触边横杠','约180×25毛坯','1','需配导向与固定限位，非承重件'),
('3mm PVC 固定后压条','约150×10～12','1','先留毛坯，按铰链让位修边'),
('2mm PVC P06活动盖','160×150','1','裸盖目标≤30g；活动总成目标≤40g、上限45g'),
], [5.7,3.4,1.0,7.8],size=9.5,center=(2,))
para('转接片、角码配合片、三角撑和局部补强均从余料加工，形状与孔位按实际轮组、舵机、铰链和安装空间确定。本章不把尚未量到的孔距作为加工尺寸。')

titles={40:'两张木板怎样排料',41:'底板与上平台怎样标孔',42:'3mm固定壳体怎样排料',43:'四壁与壳底怎样扣除板厚',44:'2mm活动盖怎样裁剪'}
notes={
40:[
'先把所有矩形完整画在两张木板上，写零件名称及尺寸，再检查是否超出原料边界。图中约5mm的边距和零件间隙是给切削与修边留的废料，不是要求锯片宽5mm。保留成品线，锯路放在废料一侧；切完再修到尺寸。',
'第一张板可把190×60立架横放在底板旁的余料带，再排80×40托片。300−220=80只是未扣边距与间隙的理论余量；应以图中的完整矩形位置为准。第二张板先切上平台，剩余区域留作转接片与三角撑。'],
41:[
'万向球轮已知安装高度20mm；若底板底面离地按参考45mm，需约25mm刚性补高，此数包含转接板、垫块及平垫的总厚度，不能另加一层后仍按25mm计算。补高后的实际水平须实测。球直径15.85mm不能代替安装高度。若改为PCB，电子板安装孔应等新PCB确定后再钻。'],
42:[
'确认实物是3mm发泡PVC，再按图划成品矩形。两侧壁和两前后壁分别标号；前后壁144mm这条边沿Y，180mm这条边竖直。操作板、横梁和触边横杠是毛坯，先不切传感器孔或开关孔。',
'薄板放在平整垫板上，用夹具固定直尺，多次浅划或使用合适细齿锯；手放在刀路外。长边先切、细条后切，窄件保持有足够夹持区，切后用细砂纸修平。不能边手托悬空板边切。'],
43:[
'两块160×180侧壁覆盖前后壁的端面。外宽150扣掉两侧各3mm，得到前后壁宽144；外长160扣掉前后各3mm，得到内部长度154，所以壳底为154×144。四壁外包络为160×150×180，上口净空为154×144。',
'先不加胶，用胶带临时拼好四壁，再试放壳底、内胆与横梁。理想外矩形对角线约219.3mm，内口约210.8mm；同一矩形两条对角线差目标不超过2mm，再检查四角。切口不方正时先修边，不靠大胶缝凑尺寸。壳底按实际内口逐边试修，以可平放取出为准，不强挤。',
'若实测板厚不是3mm，本图扣厚结果也会变化。保持外形时，前后壁宽=150−2t，壳底长=160−2t、宽=150−2t；先按实测t重画纸样，不能继续机械照抄144与154。'],
44:[
'从200×200×2板材中裁160×150活动盖，160沿车头车尾方向。修去毛刺后称裸盖，目标不超过30g；暂不钻铰链孔及侧臂孔。剩下的2mm料可做必要的轻量局部加强。',
'后固定压条使用3mm料，安装在固定壳体上，与活动盖分开。盖轴参考X=75、前缘约X=235，固定壳后缘X=70；压条参考X=64～74，实际须按铰链轴桶与叶片让位。它不随盖转动，不应挡住维护开盖或承担连杆转矩。',
'铰链、侧臂和舵机全部试摆并验证运动后，再描孔。正常开盖70°，维护需先断执行电源并拆一端拉杆。活动盖、随盖侧臂、加强片、活动铰链部分及紧固件合计实称目标≤40g、上限45g；图纸尺寸本身不能证明重量达标。']}
for k,item in enumerate(entries,1):
    n=int(item['number'])
    heading(f'62.{k}  '+titles[n],2,True)
    fn=item.get('filename',item.get('file'))
    path=HERE/item['group']/fn
    assert path.is_file(),path
    cap=re.sub(r'^图解\s*\d+\s*','',item['caption'])
    figure(path,f'图解 {n}  '+cap,item.get('alt',item['caption']),17.5)
    for text in notes[n]: para(text)

heading('62.6  实际施工顺序与复核',2,True)
for text in [
'1．对照零件表把原料分成5mm木板、3mm发泡PVC、2mm发泡PVC三堆。逐片测长宽和厚度，检查翘曲、缺角；写清材料与厚度，避免把固定壳材料拿来做活动盖。',
'2．先制作底盘、桶壳与盖的纸样。把实际电池、轮组、模块和舵机放上去，确认结构、插头方向、螺丝刀空间及内胆取出通道。',
'3．在原料上画完全部零件后再切。画成品实线、废料斜线，量两次尺寸；按图留边和间隙，确认锯缝处在废料内。每切下一件就写编号，不让余料与正式件混放。',
'4．先裁木板和PVC外形，再修边、量成品尺寸及对角线。一般外形目标误差约±1mm，嵌入壳底须按实际内口逐边试修，不以这个公差强行塞入；桶壁高度要一致。木板边去锐角，PVC切面修平，孔口和绑带槽不得有刮线毛刺。',
'5．只先加工已确定的四立柱孔。平台试装调平后，用实物转接片描电机/球轮孔，用实物电池描绑带位置。参考槽尺寸不是固定位置，任何孔槽都要避开立柱、板边薄弱处和其他螺钉。',
'6．桶壁无胶试拼、壳底试放，核对160×150外形及154×144净口。先放入内胆再确定附件位置；护线圈按所购件要求的安装孔开孔，穿线内径不是安装孔径。',
'7．试装后固定压条和活动盖，校准两铰链同轴。舵机托片先夹紧，按真实轴距50mm与高度描孔；后立柱螺钉附近预留检修开口，不把立架/托片压在不可拆的位置。',
'8．清理切屑后再装电子器件。记录切件尺寸、孔位更改、底板离地高度和活动盖总质量。完成无电机械试验，才按原手册继续单模块和整机调试。',
]:para(text)
para('合格记录：两板不翘曲，四立柱不强拉；桶壳方正且内胆可取出；盖运动无卡滞且质量达标；未知模块孔已按实物核对。此图解是设计与施工参考，仍需实际装配验证。')

assert [p.text for p in doc.paragraphs if re.match(r'步骤 \d{3}',p.text)]==steps_before
assert [[[c.text for c in r.cells] for r in t.rows] for t in doc.tables[:50]]==tables_before
assert len(doc.inline_shapes)==52
out=args.output
out.parent.mkdir(exist_ok=True,parents=True)
doc.save(out)
with zipfile.ZipFile(args.input) as z: old={hashlib.sha256(z.read(x)).hexdigest() for x in z.namelist() if x.startswith('word/media/')}
with zipfile.ZipFile(out) as z: new={hashlib.sha256(z.read(x)).hexdigest() for x in z.namelist() if x.startswith('word/media/')}
assert old<=new
audit.update({'output':out.name,'steps':330,'inline_images':52,'tables':len(doc.tables),'preserved_original_tables':50,'preserved_original_media':len(old),'new_media':len(new-old)})
(HERE/'build_audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(audit,ensure_ascii=False,indent=2))
