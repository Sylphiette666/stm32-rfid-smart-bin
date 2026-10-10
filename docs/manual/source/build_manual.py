from pathlib import Path
import json,re
from docx import Document
from docx.shared import Cm,Pt,RGBColor,Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
ROOT=Path(__file__).parent
OUT=ROOT.parent/'STM32与RFID智能移动垃圾桶_逐步组装与调试实施手册.docx'
names=['frontmatter','mechanical','electrical','software','backmatter']
packs=[json.loads((ROOT/(n+'.json')).read_text(encoding='utf-8-sig')) for n in names]
for packname,pack in zip(names,packs):
 if packname=='software':
  pack['sections']=[ss for ss in pack['sections'] if ss['title'] not in ['整机完成顺序、演示脚本和交付归档','软件故障速查与记录模板']]
  for ss in pack['sections']:
   ss['title']=ss['title'].replace('D1—D4、D6','D1至D4及D6')
   ss['intro']=ss.get('intro','').replace('用三张卡分别标记“管理员”“用户1”“用户2”；专用读写实验卡另做醒目标记。','用三张卡分别标记“管理员”“用户1”“读写实验”；第三张可兼作第二普通测试账户，但卡片块4实验不修改账户账本。需要同时保留更多用户实卡时再补卡。')
   for b in ss['blocks']:
    if b['type']=='p' and b['text'].startswith('2026-10-10 对现有'):
     b['text']='固件须按下列步骤从新工程逐模块实现。旧V2.1底板与PCF8574程序的引脚不同，不能直接烧录后接上本版MCP23017及执行器。烧录文件只有在当前接线和测试记录下通过验证，才用于整机运行。'
DOC=Document(); sec=DOC.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=Cm(1.7);sec.bottom_margin=Cm(1.7);sec.left_margin=Cm(1.8);sec.right_margin=Cm(1.8)
sec.header_distance=Cm(.65);sec.footer_distance=Cm(.7)
WIDTH=17.99
for st in DOC.styles:
 if st.type==1:
  st.font.name='微软雅黑';st.font.size=Pt(11);st.font.color.rgb=RGBColor(0,0,0)
  rp=st.element.get_or_add_rPr();r=rp.find(qn('w:rFonts'))
  if r is None:r=OxmlElement('w:rFonts');rp.append(r)
  r.set(qn('w:eastAsia'),'微软雅黑')
 for e in st.element.xpath('.//w:pBdr'):e.getparent().remove(e)
normal=DOC.styles['Normal'];normal.paragraph_format.space_after=Pt(4);normal.paragraph_format.line_spacing=1.12
normal.paragraph_format.widow_control=True
for name,size in [('Title',25),('Subtitle',13),('Heading 1',18),('Heading 2',13),('Heading 3',11)]:
 st=DOC.styles[name];st.font.size=Pt(size);st.font.bold=name!='Subtitle';st.font.italic=False
 st.paragraph_format.space_before=Pt(9);st.paragraph_format.space_after=Pt(6)
 if name.startswith('Heading'):st.paragraph_format.keep_with_next=True
DOC.styles['Heading 1'].paragraph_format.page_break_before=False
DOC.styles['Heading 3'].paragraph_format.space_before=Pt(6)
header=sec.header.paragraphs[0];header.text='STM32与RFID智能移动垃圾桶  逐步组装与调试实施手册'
for r in header.runs:r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('000000')
foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
foot.add_run('2026年10月10日  第 ')
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');foot._p.append(f)
foot.add_run(' 页')
for r in foot.runs:r.font.size=Pt(8)
def para(t='',boldlead=None,style=None):
 p=DOC.add_paragraph(style=style)
 if boldlead and t.startswith(boldlead):p.add_run(boldlead).bold=True;p.add_run(t[len(boldlead):])
 else:p.add_run(t)
 return p
def title_clean(t):
 return re.sub(r'[：:、，,/|｜（）()—–]+',' ',t).strip()
def table(headers,rows):
 n=len(headers);ratios={2:[.28,.72],3:[.27,.30,.43],4:[.18,.27,.28,.27],5:[.15,.20,.22,.22,.21]}.get(n,[1/n]*n)
 # tables of test cases require the operation cell widest
 if n==3 and any('处理' in str(h) or '通过' in str(h) for h in headers):ratios=[.23,.30,.47]
 widths=[r*WIDTH for r in ratios]
 t=DOC.add_table(rows=1,cols=n);t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
 for c,w in zip(t.columns,widths):c.width=Cm(w)
 for c,h,w in zip(t.rows[0].cells,headers,widths):c.text=str(h);c.width=Cm(w)
 for row in rows:
  cells=t.add_row().cells
  for c,v,w in zip(cells,list(row)+['']*n,widths):c.text=str(v);c.width=Cm(w)
 borders=OxmlElement('w:tblBorders')
 for side in ['top','bottom','insideH','left','right','insideV']:
  e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
 t._tbl.tblPr.append(borders)
 for ri,row in enumerate(t.rows):
  trp=row._tr.get_or_add_trPr();trp.append(OxmlElement('w:cantSplit'))
  if ri==0:trp.append(OxmlElement('w:tblHeader'))
  for c in row.cells:
   c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   mar=OxmlElement('w:tcMar')
   for side,val in [('top',65),('bottom',65),('left',75),('right',75)]:
    e=OxmlElement('w:'+side);e.set(qn('w:w'),str(val));e.set(qn('w:type'),'dxa');mar.append(e)
   c._tc.get_or_add_tcPr().append(mar)
   if ri==0:
    e=OxmlElement('w:shd');e.set(qn('w:fill'),'EEF2F4');c._tc.get_or_add_tcPr().append(e)
   for pp in c.paragraphs:
    pp.paragraph_format.space_after=Pt(2);pp.paragraph_format.space_before=Pt(1);pp.paragraph_format.line_spacing=1.05
    if ri==0:pp.paragraph_format.keep_with_next=True
    for r in pp.runs:r.font.size=Pt(10);r.bold=ri==0
 p=para();p.paragraph_format.space_after=Pt(1);p.paragraph_format.line_spacing=1
 p.add_run(' ').font.size=Pt(1)
def picture(path,caption):
 p=para();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
 r=p.add_run();r.add_picture(str(path),width=Cm(WIDTH-.2))
 for pr in r._r.xpath('.//wp:docPr'):pr.set('descr',caption)
 p=para(caption);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 for r in p.runs:r.font.size=Pt(9);r.font.color.rgb=RGBColor.from_string('555555')
para('STM32与RFID\n积分式智能移动垃圾桶','Title' if False else None,style='Title')
para('逐步组装与调试实施手册',style='Subtitle')
para('单台制作  复用现有模块  全部33项中等及以下功能')
para('装配版本 2026年10月10日')
para('从配件核对、板材加工、机构试装到逐根接线、程序实现、整机验收，本册按制作顺序给出可执行步骤。先检查再上电，先做单项测试再合装，每个阶段均保留实测记录。')
table(['本册采用','制作时关注'],[
 ['核心板加洞洞板','STM32F103C8T6与MCP23017扩展；不用旧版PCF8574底板的接线。'],
 ['轻型机械结构','3mm PVC固定外壳，2mm PVC活动盖，1mm PP五片内胆。'],
 ['一台完整样机','RFID积分、轻型开盖、循迹停站、蓝牙、显示、语音与环境检测。'],
 ['分阶段完成','机械加工和软件驱动可交叉推进；电源及停止回路通过后才接执行器。']])
para('请先在到货记录中确认模块型号、引脚、轮组孔距和急停接点条件；按本册的检查标准测量后填写实际参数。章节末有故障排查、制作记录和演示顺序，可用于课程答辩准备。')
tocp=para('目录');tocp.paragraph_format.page_break_before=True
for rr in tocp.runs:rr.font.size=Pt(18);rr.bold=True
toc=DOC.add_paragraph();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'TOC \\o "1-1" \\h \\z \\u');toc._p.append(fld)
para('可点击目录或使用Word导航窗格跳转。步骤编号全册连续，便于在接线检查和问题记录中引用。')
stepn=0;sectionn=0;stats=[];source=[]
figmaps=json.loads((ROOT/'figure_map.json').read_text(encoding='utf-8')) if (ROOT/'figure_map.json').exists() else {}
for name,pack in zip(names,packs):
 count=0
 for si,ss in enumerate(pack['sections']):
  sectionn+=1
  rawtitle=ss['title'];title=title_clean(rawtitle)
  chapter=DOC.add_heading(f'{sectionn:02d}  {title}',1)
  chapter.paragraph_format.page_break_before=(si==0 and name!='backmatter')
  if ss.get('intro'):para(ss['intro'])
  if rawtitle in figmaps:
   for fig in figmaps[rawtitle]:picture(ROOT/fig['path'],fig['caption'])
  for b in ss['blocks']:
   typ=b['type']
   if typ=='p':para(b['text'])
   elif typ=='table':table(b['headers'],b['rows'])
   elif typ=='steps':
    for item in b['items']:
     stepn+=1;count+=1
     DOC.add_heading(f'步骤 {stepn:03d}  '+title_clean(item['title']),3)
     group=[para(item['action']),para('□ 检查  '+item['check'],boldlead='□ 检查')]
     if item.get('fix'):group.append(para('未通过时  '+item['fix'],boldlead='未通过时'))
     for i,p in enumerate(group):
      p.paragraph_format.keep_with_next=(i<len(group)-1)
      p.paragraph_format.keep_together=True
   elif typ in ['code','pre']:
    p=para(b.get('text',b.get('code','')))
    for r in p.runs:r.font.name='Consolas';r.font.size=Pt(9)
   elif typ in ['heading','h2']:DOC.add_heading(title_clean(b.get('text',b.get('title',''))),2)
   elif typ=='image':picture(ROOT/b['path'],b.get('caption','装配示意'))
   else:raise RuntimeError((name,typ))
 stats.append({'part':name,'sections':len(pack['sections']),'steps':count})
 source+=pack.get('sources',[])
sectionn+=1;DOC.add_heading(f'{sectionn:02d}  技术资料与参数依据',1)
para('芯片和模块的电气条件应以所购型号的原厂资料为准。下面列出核对本册设计使用的资料。机械外形、程序阈值、动作顺序和测量合格线是本项目的设计设定，须按各章工步进行样机验证。资料核对日期为2026年10月10日。')
grouped={}
for ss in source:
 if isinstance(ss,str):ss={'title':ss,'url':ss}
 url=ss.get('url','');title=ss.get('title',ss.get('name','技术资料'))
 if url:
  entry=grouped.setdefault(url,{'titles':[],'notes':[]})
  if title not in entry['titles']:entry['titles'].append(title)
  if ss.get('note'):entry['notes'].append(ss['note'])
seen=set(grouped)
for url,entry in grouped.items():
 p=para('；'.join(entry['titles']));p.paragraph_format.keep_with_next=True
 for r in p.runs:r.bold=True;r.font.size=Pt(10)
 for note in entry['notes']:
  p=para(note);p.paragraph_format.keep_with_next=True
 p=para(url)
 for r in p.runs:r.font.size=Pt(9)
DOC.core_properties.title='STM32与RFID智能移动垃圾桶逐步组装与调试实施手册'
DOC.core_properties.subject='全部33项中等及以下功能的配件装配接线程序实施与验收'
DOC.core_properties.author='';DOC.core_properties.keywords='STM32 RFID 逐步组装 接线 机械 调试 33项功能'
DOC.save(OUT)
(ROOT/'build_audit.json').write_text(json.dumps({'output':str(OUT),'sections':sectionn,'steps':stepn,'parts':stats,'source_count':len(seen),'paragraphs':len(DOC.paragraphs),'tables':len(DOC.tables),'figures':len(DOC.inline_shapes)},ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'output':str(OUT),'steps':stepn,'sections':sectionn},ensure_ascii=False))
