from pathlib import Path
from collections import Counter, defaultdict
import argparse, csv, hashlib, json, re, zipfile
from docx import Document
from PIL import Image

HERE=Path(__file__).resolve().parent
ROOT=HERE.parents[1]
NAME='STM32与RFID智能移动垃圾桶_逐步组装与调试实施手册_图解增强版.docx'
ap=argparse.ArgumentParser();ap.add_argument('--render',default='render_final');ap.add_argument('--allow-pending',action='store_true');args=ap.parse_args()
source=Document(HERE/'before.docx'); dest=Document(HERE/'staging'/NAME)
steps=lambda d:[p.text for p in d.paragraphs if re.match(r'步骤 \d{3}',p.text)]
def media(path):
    with zipfile.ZipFile(path) as z:
        return Counter(hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/'))
oldmedia=media(HERE/'before.docx'); newmedia=media(HERE/'staging'/NAME)
figure_status=[]
for filename in ('45_pcb_top.png','46_pcb_bottom.png','47_power.png','48_modules.png','49_conditions.png'):
    path=HERE.parent/'a-design'/'diagrams'/filename
    sha=hashlib.sha256(path.read_bytes()).hexdigest() if path.exists() else None
    figure_status.append({'file':filename,'source_sha256':sha,'latest_embedded':bool(sha and newmedia[sha])})
body='\n'.join(p.text for p in dest.paragraphs)
interfaces=[]; interface_values=[]
for t in dest.tables:
    if t.cell(0,0).text=='板侧针号':
        interfaces.extend(r.cells[0].text for r in t.rows[1:])
        interface_values.extend(tuple(c.text for c in r.cells[:3]) for r in t.rows[1:])
with (HERE.parent/'a-design'/'connector-map.csv').open(encoding='utf-8-sig',newline='') as f:
    expected_interfaces={(f"{r['Ref']}.{r['PCB_pin']}",r['Net'],r['External_signal_by_label']) for r in csv.DictReader(f)}
build=json.loads((HERE/'build_audit.json').read_text(encoding='utf-8'))
expected_tables=json.loads((HERE/'baseline_tables.json').read_text(encoding='utf-8'))
for change in build['changed_cells']:
    expected_tables[change['table']][change['row']][change['col']]=change['after']
actual_tables=[[[c.text for c in r.cells] for r in t.rows] for t in dest.tables[:len(expected_tables)]]
report={'steps_preserved':steps(source)==steps(dest),'steps':len(steps(dest)),
        'original_media_preserved':all(newmedia[k]>=v for k,v in oldmedia.items()),
        'original_media_count':sum(oldmedia.values()),'new_media_count':sum(newmedia.values()),
        'new_figures_latest':all(f['latest_embedded'] for f in figure_status),'figure_sources':figure_status,
        'interfaces':len(interfaces),'unique_interfaces':len(set(interfaces)),
        'interface_values_match_csv':set(interface_values)==expected_interfaces,
        'original_tables_only_authorized_changes':actual_tables==expected_tables,
        'has_placeholders':'【待最终图解' in body,'inline_images':len(dest.inline_shapes),
        'tables':len(dest.tables),'sha256':hashlib.sha256((HERE/'staging'/NAME).read_bytes()).hexdigest(),
        'original_still_unchanged':hashlib.sha256((ROOT/NAME).read_bytes()).hexdigest()=='f7d0c5ce51f904e58d40f9733fe5b030d29f31db53e26faf53ad6d21b3713187'}
if not args.allow_pending:
    assert not report['has_placeholders']
    assert report['new_figures_latest'], figure_status
    assert report['inline_images']==57
assert report['steps_preserved'] and report['steps']==330 and report['interfaces']==181 and report['unique_interfaces']==181
assert report['original_media_preserved']
assert report['interface_values_match_csv'] and report['original_tables_only_authorized_changes']
reference=ROOT/'.cutting-diagrams-20261010'/'render_final';current=HERE/args.render
def crop_hash(p):
    im=Image.open(p).convert('RGB')
    return hashlib.sha256(im.crop((25,70,im.width-25,im.height-70)).tobytes()).hexdigest()
if current.exists():
    hashes=defaultdict(list)
    for p in reference.glob('page-*.png'): hashes[crop_hash(p)].append(int(p.stem.split('-')[1]))
    matched=[];changed=[]
    pages=sorted(current.glob('page-*.png'),key=lambda p:int(p.stem.split('-')[1]))
    for p in pages:
        n=int(p.stem.split('-')[1]); h=crop_hash(p)
        if h in hashes: matched.append({'current':n,'prior':hashes[h]})
        else: changed.append(n)
    report['render_pages']=len(pages);report['unchanged_body_pages']=matched;report['changed_or_new_pages']=changed
(HERE/f'audit_{args.render}.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in report.items() if k!='unchanged_body_pages'},ensure_ascii=False,indent=2))
