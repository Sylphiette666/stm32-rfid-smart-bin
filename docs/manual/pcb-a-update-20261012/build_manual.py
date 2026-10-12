from pathlib import Path
from copy import deepcopy
import argparse, csv, hashlib, json, re
from collections import defaultdict
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE=Path(__file__).resolve().parent
ROOT=HERE.parents[1]
DESIGN=HERE.parent/'a-design'
NAME='STM32与RFID智能移动垃圾桶_逐步组装与调试实施手册_图解增强版.docx'
BASE_SHA='f7d0c5ce51f904e58d40f9733fe5b030d29f31db53e26faf53ad6d21b3713187'
ap=argparse.ArgumentParser(); ap.add_argument('--final',action='store_true'); args=ap.parse_args()
assert hashlib.sha256((HERE/'before.docx').read_bytes()).hexdigest()==BASE_SHA
doc=Document(HERE/'before.docx')
ps=list(doc.paragraphs); ts=list(doc.tables)
steps_before=[p.text for p in ps if re.match(r'步骤 \d{3}',p.text)]
audit={'baseline_sha256':BASE_SHA,'changed_paragraphs':[],'changed_cells':[],'figures':[],'status':'staged final pending rendered QA and root release' if args.final else 'text draft; layout/production not yet released'}
exec(compile((ROOT/'.cutting-diagrams-20261010'/'retained_formatting.py').read_text(encoding='utf-8'),'<retained-formatting>','exec'))
def patch(i,new):
    p=ps[i]; old=p.text
    assert not old.startswith('步骤 '), (i,old)
    replace_text(p,new); audit['changed_paragraphs'].append({'index':i,'before':old,'after':new})
def cell(t,r,c,new):
    q=ts[t].cell(r,c); old=q.text
    p=q.paragraphs[0]; replace_text(p,new)
    for extra in list(q.paragraphs)[1:]: extra._element.getparent().remove(extra._element)
    audit['changed_cells'].append({'table':t,'row':r,'col':c,'before':old,'after':new})
def ins_after(index,text):
    p=para(text)
    ps[index]._p.addnext(p._p)
    audit.setdefault('inserted_after',[]).append({'index':index,'text':text})

replacements={
3:'装配版本 2026年10月12日  PCB A1',
4:'保留原版330个步骤与33项功能，现含49幅操作图解及1幅整机外观概念图；原有图1至图7继续保留。当前按SmartBin_A A1裸PCB手焊方案施工。',
14:'当前电气施工先读第63章图解45—49及接口表，再按原330步骤逐项调试。图解01—30说明操作原理，31—35为SG90连接详图，40—44为木板与PVC裁剪图。第61章图解36—39与142条洞洞板焊线完整保留为旧版归档，不用于本次PCB焊接。模块均按实物信号名称逐线连接。',
117:'图解 03  旧洞洞板焊接方法示例。PCB A1按第63章焊接；焊点发亮程度不能单独证明接触可靠。',
118:'本图保留用于识别焊盘与焊接基础。当前使用SmartBin_A A1双层PCB替代原PWR、IO两块洞洞板，板内连接由铜箔完成；不得再照旧孔位补142条飞线或两根裸汇流线。模块仍在板外固定并用编号线束连接。',
261:'用已有10mm双通柱制作电子器件安装分区。先用纸板摆放核心、100×100mm PCB、MCP模块、驱动和电源模块，留插拔线与螺丝刀空间，再按实物描安装孔；托板外形可按实际空间裁定。',
262:'机械托板数量不等于电路板数量。本机以一块SmartBin_A A1 PCB替代旧PWR与IO；核心、MCP23017成品模块、TB6612及电源模块仍独立固定，不叠压在裸露焊点上。',
500:'六组KF301端子焊在PCB的J1—J6处，不悬空吊在线上。夹紧导线后轻拉检查，端子外不得露出散铜丝。XT30U接头外壳固定或设短服务线，杜邦信号接头加适度线夹。',
541:'本章逐根接线采用STM32F103C8T6可拆核心板、SmartBin_A A1 PCB和MCP23017成品模块。PA/PB是STM32端口，GPA/GPB是MCP端口，例如PB0与GPB0不是同一个脚。以下保留原步骤编号；PCB板侧引脚以第63章为准，不能把表格顺序当成商品模块的物理针序。历史PCF8574 V2.1板及第61章洞洞板孔号均不适用于本次施工。',
542:'本PCB为100×100mm双层板，集中原电源分配、ADC、照明电阻、AHCT缓冲、三管急停和信号汇总。核心、MCP、驱动、降压/稳压和传感器外置；需5V数字输入适配时仍在模块侧处理。按新位号装件，原PWR/IO位号只保留在归档对照中。',
545:'把配件分为核心/通信、传感器、执行器、电源/接线、备用五盒。先找出标有SmartBin_A A1的PCB；按第63章标记各线束两端，如J20.1 SERVO_SIG、J15.3 CS、3V3_EXT。记录每个模块丝印、型号、供电、输出类型和资料来源。',
557:'先在裸PCB上空插排针、端子、IC座、电阻、电容，核脚距与外壳是否互相挤压；按第63章从低件到高件焊。排针先焊一脚、校直，再焊其余脚。柔性线剥约4—5mm，端子压线处按其说明压接端头或使用整齐裸铜，不把厚焊锡堆夹在螺丝下。',
566:'21  PCB电源区 电池 保险 开关 降压与去耦',
586:'断电。外置3.3V模块IN+接J7.1（5V_BUS），GND/IN-/OUT-接J7.2，OUT+接J7.3（3V3_EXT）。空载测3.3V，再用小负载分级检查预期电流和温升。该模块与核心3V3输出不并联；至少200mA仅是原采购下限，全部外设负载加20%余量不得超过实测持续能力。',
591:'PCB上C1=470μF对应电机支路，C2=1000μF对应舵机支路，C3=100μF对应逻辑支路；C4/C5/C6分别为100nF，C8在AHCT附近。电解1脚为正、2脚为GND；板位预留5mm脚距、直径不大于12.5mm。实物不合时用短绝缘引线并独立固定。远端MCP、驱动与DFPlayer的就近去耦仍要保留。',
596:'PCB的J1接开关后电池，J2接降压输入，J3接已调好的5V降压输出，J4/J5/J6分别分给电机、舵机、逻辑。5V_BUS是同一电气网络，布线按三支路分流；旧版两根0.8mm汇流线及W001—W007不再焊。电池/降压主干用20AWG短线，电机/舵机短支路22AWG，压降偏大时改20AWG。主电流不得经信号排针或普通杜邦线。',
603:'3A保险管是初选值，熔断曲线和直流额定值须与实测启动电流、线材及电池保护板协调，保护板与保险不能相互省略。SG90先按2A瞬态留量并实测，不能把0.8A估算当上限；控制上禁止轮驱动与开盖并发。两轮各1A加其余0.5A预算，移动模式约2.5A；舵机2A加其余0.5A，开盖模式也约2.5A。持续3A、短时至少4A是分时正常工况的选型起点，不代表支持电机和舵机同时堵转的约4.5A。语音音量、模块和机械负载会改变电流，最终以实际压降、温升及峰值测量确认。',
612:'J8是本项目自定义17针核心PA线束，J9是16针PB线束，不是核心原有整排插座。按第63章逐根核对：J8.1供核心5V、J8.2地、J8.3接核心3V3参考；PA1来自板上ADC经JP9。SWD仍直接接核心PA13/PA14。不要整排直插或把J8.3与3V3_EXT相连。',
630:'给核心USB贴“本机不接USB，PA11/12已分配”标签。正常系统下载用SWD、通信接J18三线USB串口，转换器VCC全部不接。独立USB台架须先断本板电源，并拔下与核心相连的全部线束后另测；只拔5V线并不能隔离ADC和其他GPIO反灌路径。',
632:'未通过时  需要核心USB测试时先彻底移出当前系统线束；不能在PA11/PA12仍连接业务模块时直接插核心USB。',
635:'23  PCB总线区 MCP23017 I²C 上拉与OLED',
637:'已焊好的MCP模块由J10接3V3_EXT、GND、SCL、SDA、RESET五线；另由J11接GPA0—GPA6七线、J12接GPB0—GPB6七线，共19根线。板内再分到对应功能接口。A0/A1/A2在模块侧接地；INTA/INTB不接；GPA7/GPB7不引出且程序设输出0。RESET不接核心NRST。',
682:'RC522使用J15：1=3V3_EXT、2=GND、3=模块SDA/SS（片选PA4）、4=SCK、5=MOSI、6=MISO、7=RST（GPA4），IRQ不接。板上R11为片选10k上拉；模块商品排针次序不同，要逐线对应，不把SDA接I²C。',
694:'抄下FLASH芯片全名并核为3.3V SPI NOR。J16的1/2为3V3/GND，3=CS（PA15），4=CLK，5=DI/SI，6=DO/SO。7/8为WP#/HOLD#，仅实物确有对应外露脚时使用；R13/R14按原有上拉决定装，R12是CS上拉。1.8V芯片不能直接接入。',
729:'确认编码器模块能用3.3V后按J24接VCC/GND/CLK/DT/SW。CLK/DT只选一组上拉：移除模块原上拉后装R27/R28各10k拉3V3_CORE；或保留已确认拉3V3_EXT的模块上拉并不装R27/R28。SW按R33、模块同轨上拉或已验证MCP内部上拉三者确定有效上拉。不得让同一信号分别拉到两个稳压轨。',
739:'MODE与BACK使用常开轻触按钮。J25.1—2接MODE，3—4接BACK；按下才跨组导通。R34/R35默认已装10k上拉3V3_EXT。J25.5—6留给闭盖NO微动，不能把三组开关的信号脚短接。',
763:'在PCB定位Q1/Q2/Q3，分别对应原步骤中的QES1/QES2/QES3，不另加一组三管。按实际厂家资料核E/B/C：PCB1=E、2=B、3=C。三个E接地，三个B均有100k下拉；断电核对导通及阻值后才继续。',
764:'TO-92外壳平面方向不能代替脚号。若所购品牌脚序与板定义不同，应按电气脚名弯脚、套绝缘管并固定，避免交叉脚碰触；不要直接按图形朝向硬插。',
792:'按下急停使TB6612待机，输出高阻，车轮可能惯性滑行；它不是机械刹车，也不切断舵机电源。舵机停止由软件处理，整机断电靠总开关。三管钳位的Q2偏置依赖3V3_EXT；EXT掉电而核心仍输出高时不能保证完整硬件联锁。正常电源同启同停，异常用总开关断电；本课设不声明工业安全等级。',
798:'□ 检查  IC座14脚接5V、7脚接GND；1/4/10使能脚均接GND，13脚接5V。第3路9脚接PB9、8脚接BUZZ_5V，第4路12脚接地、11脚悬空；电源无短路。',
801:'断电插U1，暂不接执行器。分别令PA0、PB15低/高，检查3脚、6脚0/5V；再测PWM与10μs脉冲。第3路PB9已有R36或R37唯一偏置，检查9脚与8脚同相跟随。仅第4路禁用，不能套旧图把9脚接地、10脚接5V。',
805:'PCB已将Y1接J20.1舵机PWM，Y2接J21.3超声TRIG，Y3接JP3.3。J20.1是5V逻辑信号而非舵机正电源；SG90供电走J5。禁止另跨线把PA0/PB15直接并到相应5V输出。',
839:'SG90红线从J5.1取5V，棕/黑线直接回J5.2；信号线经实物核对接J20.1，即U1第3脚Y1。J20.2仅为信号参考地，不承担舵机主电流。先拆杆以50Hz、1500μs空载检查，再按第60章机械图装连杆。',
894:'PCB已焊R25=12k上臂、R26=20k下臂：J21.4原始ECHO→R25→PB1，PB1经R26到GND。先测分压再接核心；不另跨原始ECHO到PB1，也不在该节点并大电容以免损失脉宽。',
906:'支持3.3V的DHT版本接J22的电源、GND、DATA；JP7短1—2，JP8闭合。PB11开漏，R29或模块已有同轨上拉须有效。裸四脚DHT按厂家引脚图核对，第三脚NC不接。未确认电压前JP7/JP8保持开路。',
910:'必须5V的DHT：JP7短2—3，JP8始终开路；外接BSS138双向开漏转换板到J23，1=LV/3V3_EXT、2=HV/5V、3=GND、4=低压通道PB11、5=高压通道DHT_DATA_RAW。确认两侧有效上拉；不能用单向NPN代替，也不能短JP8旁路转换。',
926:'AO3400A仍是已购SOT-23裸管，先焊到已购SOT-23转DIP板，再用标明G/S/D的线束接PCB J28。转接板独立绝缘固定，不假定商品转接板排针顺序。它只驱动一颗小白LED，不驱动电机或舵机。',
931:'用镊子对准已少量预上锡的转接板焊盘，固定AO3400A一脚，核对后焊其余两脚。断电查无连锡并追踪真实G/S/D。J28.1→G、J28.2→S/GND、J28.3→D/LED_NEG。转接板独立固定并绝缘，不整排猜接。',
936:'PCB已放R3=100Ω栅极串联、R4=100k下拉、R5=220Ω限流。GPA5经J11.6进入板内；MOS接J28，白LED接J29（1正2负）。无需旧PWR的JP10/JP11/JP12线束或另焊一组相同电阻。',
948:'PCB已固定AHCT第3通道9脚接PB9、10脚接地。3.3V可控制的蜂鸣模块将JP3短1—2；确需5V控制的短2—3，不允许两只帽并接两个输出。JP4按模块电源选择。高有效默认装R36=10k下拉、R37不装；低有效先拆R36再装R37拉3V3_CORE，即使不用蜂鸣也须保留其中一只偏置。第61章“五拆六增”只属于旧洞洞板，PCB不拆铜箔、不照做。',
990:'PCB上R1=33k、R2=10k、C7=100nF构成电池分压；JP9闭合后送J8.5的PA1。8.4V理论约1.953V，长采样时间并按实际Vref校准。单独拔核心电源前先断全机电源；若电池仍在，至少先拔JP9隔离ADC，但其他GPIO仍可能反灌，不能把JP9当全接口隔离。',
1524:'PCB用四个M3安装孔和10mm绝缘立柱固定，孔中心相对100×100mm板左上为(5,5)、(95,5)、(5,95)、(95,95)。先把实物放到底层，与电池、电源模块和轮线一起试排，再描孔，端子与螺丝刀方向可达。',
1525:'焊接面与木板、螺钉及线束之间保留可靠间隙和绝缘；切除过长引脚，不能让裸焊点接触金属件。旧版汇流铜线不装在PCB上。量实际板孔再配钻，板底不得残留切屑。',
1759:'61  旧版洞洞板归档 不用于PCB施工',
1760:'本章完整保留2026年10月10日的PWR/IO洞洞板参考布局R1及142条板内连线，用于历史追溯。用户现已选择SmartBin_A A1 PCB，请跳到第63章施工。本章A—X孔号、旧位号、飞线、裸汇流线及AHCT五拆六增均不适用于新PCB，不得混合执行。',
1773:'图解 36  旧版归档 PWR洞洞板元件面。仅用于历史对照，不用于PCB A1施工。',
1776:'图解 37  旧版归档 PWR洞洞板焊接面。此图已镜像，其铜线施工不用于PCB A1。',
1779:'图解 38  旧版归档 IO洞洞板元件面。旧AHCT第3通道配置不用于PCB A1。',
1782:'图解 39  旧版归档 IO洞洞板焊接面。旧孔号、飞线与新PCB位号不可互换。',
1792:'以下内容仅是旧洞洞板条件改线记录，不用于PCB A1。新板AHCT第3通道已经永久接PB9并使能，只按第63章设置JP3和R36/R37。不得在新PCB上执行以下五拆六增。',
1889:'2026年10月12日归档：旧图、孔位表及142条焊线不用于PCB施工。A1位号、AHCT接法和制造文件见第63章；设计检查不代替实物验收。',
}
for i,v in replacements.items(): patch(i,v)
audit['removed_paragraphs']=[{'index':1890,'before':ps[1890].text,'reason':'将归档日期及验收边界并入前一段，避免独立尾段占整页'}]
ps[1890]._element.getparent().remove(ps[1890]._element)
ins_after(87,'PCB A1施工顺序：先看第63章确认板号、装件和跳帽，再回到步骤091—194逐模块检查，最后执行软件、整机和33功能验收。第61章为归档，跳过其中旧飞线施工。')
for i in range(1761,1891):
    if ps[i].style.name.startswith('Heading') and '旧版归档' not in ps[i].text:
        patch(i,ps[i].text+'  旧版归档')
for section in doc.sections:
    for footer in (section.footer,section.first_page_footer,section.even_page_footer):
        for paragraph in footer.paragraphs:
            for run in paragraph.runs:
                if '2026年10月10日' in run.text: run.text=run.text.replace('2026年10月10日','2026年10月12日')
                if '2026 年 10 月 10 日' in run.text: run.text=run.text.replace('2026 年 10 月 10 日','2026 年 10 月 12 日')

cell(0,1,0,'核心板加PCB A1'); cell(0,1,1,'STM32F103C8T6与MCP23017成品模块；裸PCB手焊，旧PCF8574板与洞洞板孔号不适用。')
cell(1,1,1,'核心板拆出独立固定；SmartBin_A A1 PCB汇总接线，模块在板外按信号名用编号线束连接。')
cell(3,3,2,'焊PCB直插件、AO转接板、接线端子和电容。')
cell(5,21,2,'AO裸管先焊至已购转接板，再按实际G/S/D接J28并绝缘固定；不是商品板直插针序。')
cell(5,23,2,'U1将舵机、超声和可选蜂鸣控制缓冲为5V；配DIP14座，C8为100nF。不能换成74HC125。')
cell(5,33,2,'J1—J6共6个；5.08mm脚距，需核实外壳及脚形与PCB封装，螺丝刀与进线空间可达。')
cell(5,34,0,'SmartBin_A A1 裸PCB 双层100×100mm'); cell(5,34,1,'实际装配1块 订板数量按订单'); cell(5,34,2,'1.6mm、1oz，已有洞洞板保留备用；按新制造文件订裸板，不下SMT贴片单。')
cell(5,35,1,'本板信号/跳帽共169针，至少5条40Pin排针'); cell(5,35,2,'从已购2.54mm单排排针按接口裁切，不足才补；模块自身排针及损耗另计，跳帽另备。')
cell(5,37,1,'已有保留备用'); cell(5,37,2,'PCB A1不用旧洞洞板两根裸汇流线，不再为本板补购；主电流仍用规定绝缘导线。')
cell(13,5,2,'J1—J6采用5.08mm直插件；先核孔径、脚形和外壳，螺丝一字或十字不改变电气功能。')
cell(18,7,0,'PCB接口 端子及标签'); cell(23,5,2,'电机/舵机/电池主电流不通过它；验证后接入PCB A1对应接口。')
cell(18,0,0,'本电路用件 已计入PCB清单')
cell(18,1,2,'PCB Q1/Q2/Q3对应原QES1/QES2/QES3，不重复增购；核实所购品牌E/B/C。')
cell(18,2,2,'PCB R15/R17为NC1/NC2基极限流，R19为Q2基极上拉。')
cell(18,3,2,'PCB R16/R18/R20为三个基极的100k下拉。')
cell(18,4,2,'PCB R21串PB5到STBY，不可用导线旁路。')
cell(18,5,2,'PCB R22/R23下拉PB5/STBY，R24拉PB13到CORE。')
cell(22,4,1,'J30.1/2/3接VCC/GND/IN；JP4选电源，JP3选直出或5V缓冲。')
for ri,row in enumerate(ts[20].rows):
    if ri: audit.setdefault('ahct_table_before',[]).append([c.text for c in row.cells])
ahct_rows=[('14 VCC / 7 GND','5V_BUS / GND；C8=100nF','5V供电；DIP14座'),('1 / 4 / 10使能','全部接GND','通道1、2、3永久使能'),('2 A1 / 3 Y1','PA0 / SERVO_SIG_5V','J20.1仅信号；舵机电源走J5'),('5 A2 / 6 Y2','PB15 / US_TRIG_5V','J21.3超声TRIG'),('9 A3 / 8 Y3','PB9 / BUZZ_5V','JP3选直出或缓冲；R36/R37只装一只'),('12 A4 / 13使能 / 11 Y4','GND / 5V_BUS / 悬空','第4路禁用；11脚不接'),('R6 / R7 / R36或R37','输入默认低或蜂鸣不响电平','不使用蜂鸣也不可让PB9输入浮空')]
for ri,row in enumerate(ahct_rows,1):
    for ci,v in enumerate(row): cell(20,ri,ci,v)
for ri,vals in enumerate([
('红色电源线','J5.1 5V_BUS','来自降压5V；不接8.4V或3.3V'),
('棕色或黑色地线','直接回J5.2 GND','承载舵机电流，不经信号地回流'),
('橙 黄 白信号线','J20.1 SERVO_SIG_5V','5V逻辑PWM，来自U1.3；不是正电源'),
('核心PA0','J8.4 → U1.2 A1','经缓冲后去SG90；全系统共地')],1):
    for ci,v in enumerate(vals): cell(44,ri,ci,v)
for ti in (45,46,47,48,49):
    cell(ti,0,0,'旧版归档 '+ts[ti].cell(0,0).text)

circuit=json.loads((DESIGN/'circuit.json').read_text(encoding='utf-8'))
parts=circuit['parts']; byref={p['ref']:p for p in parts}
with (DESIGN/'connector-map.csv').open(encoding='utf-8-sig',newline='') as f: connectors=list(csv.DictReader(f))
assert len(connectors)==181 and len(parts)==94
audit['circuit_sha256']=hashlib.sha256((DESIGN/'circuit.json').read_bytes()).hexdigest()
def numbered(items):
    for n,text in enumerate(items,1): para(f'{n}．{text}')
def fig(number,filename,caption,alt):
    path=DESIGN/'diagrams'/filename
    if not path.exists():
        if args.final: raise FileNotFoundError(path)
        para(f'【待最终图解 {number}：{caption}】')
        audit.setdefault('pending_figures',[]).append(str(path))
    else:
        figure(path,f'图解 {number}  {caption}',alt,17.5)
        audit['figures'][-1]['sha256']=hashlib.sha256(path.read_bytes()).hexdigest()

heading('63  PCB A1手焊装配与模块连接',1,True)
para('本章是当前施工入口。用一块SmartBin_A A1裸PCB替代原PWR与IO两块洞洞板，复用已购买的直插件和成品模块，保留原33项功能与330步骤的调试顺序。板内连接由PCB铜箔完成，模块之间仍可使用短杜邦信号线。')
para('本次设计包含37个电阻焊位、8个电容、3只2N3904、1只AHCT、33组外接接口、8组跳帽/测试接口和4个安装孔。条件焊位不一定全部装，R36/R37互斥。板上没有STM32、MCP、RFID或降压芯片；这些模块继续在板外。')
para('第61章及图解36—39保留为旧版归档，不能照其A—X孔号补飞线，也不能执行五拆六增。本板U1第3通道已永久接PB9并使能，用JP3选择蜂鸣控制电平。机械部分继续用第60、62章。')
para('原步骤中的5V_MOTOR、5V_SERVO、5V_LOGIC表示三条用途不同的供电支路，在A1电气网络中都属于5V_BUS。它们必须按J4/J5/J6分开接线，不能把“同网”理解成电机电流可穿过任意小信号线。')
heading('63.1  先辨认板子与已有配件',2)
para('PCB规格为100×100mm、双层、1.6mm厚、1oz铜。四个M3孔为3.2mm非金属化孔；以元件面左上板角为原点，孔中心为(5,5)、(95,5)、(5,95)、(95,95)mm。按板上A1版号和真实制造文件核对，不能用旧V2.1的Gerber。')
fig(45,'45_pcb_top.png','PCB元件面装配总览','SmartBin_A A1元件面，按实际位号装件及辨认接口一脚')
para('元件从器件面插入、在另一面焊接。翻板后左右改变，但编号仍指同一物理孔；正反图不可再次自行镜像。板上“PIN 1=SQUARE PAD”只用于J/JP接口；电解电容必须按“＋”确认正脚，不凭轮廓形状判断极性。')
fig(46,'46_pcb_bottom.png','PCB焊接面与翻板方向','SmartBin_A A1从焊接面观察的实际PCB，接口编号与元件面相同')
para('先拿实物空插，不急着焊。轴向电阻脚距10.16mm、目标本体不大于6.3×2.5mm；DIP14插座排距7.62mm；电解预留5mm脚距、最大直径12.5mm；瓷片电容5mm脚距。J1—J6采用Phoenix库的2P端子几何：孔距5.08mm、钻孔1.3mm、包络约11.22×11.125mm。复用已有KF301，核实引脚能穿孔、相邻外壳不挤、进线朝板边且螺丝刀可操作；不要求重买Phoenix品牌。')
heading('裸PCB怎样下单',3)
para('上传本次A1制造包Gerber_RS274X.zip，按100×100mm、2层、1.6mm、1oz、普通FR-4喷锡板核对。只做裸PCB，不选SMT/PCBA或钢网，不需要提交贴片坐标CPL。BOM留作自己装件核对，不把它作为自动代贴采购单。')
para('只装配一台；实际起订片数和优惠以裸板下单界面为准，不受先前SMT页面“至少2片”限制。平台若按多片一批制作，多余裸板可留备份，不必给每块板都买元件。先看生产预览：尺寸、闭合板框、4个安装孔、排针/端子孔和正负极标记均与本次A1一致，再确认报价。')
para('本版规则：最小线距0.20mm，信号线通常0.30mm、局部0.25mm；常规过孔外径0.8mm/钻孔0.4mm，铜到板边至少0.5mm。电源铜宽由实际布局单独决定，不能用信号线宽概括。修改板图后须重新运行DRC并重新导出制造包。')
heading('已有配件与少量补料',3)
table(['物品','本板需求','处理方法'],[
('电阻 电容 三极管 AHCT','37R焊位、8C、3Q、1U','优先从原采购包分拣，条件位按63.3决定'),
('DIP14座','1只','BOM中U1_SOCKET为附件行；先焊座，再插1颗U1，不是两颗芯片'),
('KF301-2P 5.08mm','6只','用原已购件；脚形或外壳不符先处理'),
('2.54mm单排针','本板169针','至少5条40Pin，可剪短；已有不足才补'),
('2.54mm跳帽','运行最多约7只','建议备10只；JP10正常不装'),
('裸PCB','装配1块','下单数量按平台；多余板留备份，不重复买模块'),
('AO3400A与SOT23转接板','原各1只/块','单独焊好，核G/S/D后接J28'),
('洞洞板 0.8mm裸铜线','本板不用','已买留备用，不再焊旧汇流线'),
], [4.5,4.1,9.3],size=9.5)
para('初步新增预算仍按裸PCB、自焊配件补缺及普通物流约20—60元考虑，不含焊接工具和已经买好的模块；这不是下单实时报价。尺寸、数量、运费和优惠以实际制造订单为准。')

heading('63.2  裸板到手后的焊接顺序',2,True)
numbered([
'准备烙铁架、镊子、斜口钳、助焊剂、焊锡、万用表和放大镜。桌面放绝缘垫，板上不接电池、USB或外置模块。把U1芯片单独收好，先不插座。',
'对照位号分拣电阻，把33k、12k、20k与其余阻值分开；用万用表确认。R1/R2电池分压和R25/R26超声分压采用1%精度。把待定/DNP电阻另放一格，避免先全焊再返工。',
'先装固定电阻。弯脚处离本体留余量，不扯根部；从元件面插入，翻面固定。加热焊盘和引脚后送少量锡，焊锡润湿两者；冷却后剪去多余脚，不能让残脚碰邻焊盘。',
'装C4—C8五只100nF瓷片电容，无正反。C7属于ADC滤波，C8属于U1去耦，按位号放置；不要因容量相同漏掉其中一只。',
'插DIP14座，缺口与丝印一致。先焊对角两脚，回元件面看座是否贴平，校正后焊余脚。再查相邻14脚无锡桥，芯片仍不插。',
'装Q1/Q2/Q3。依据所购品牌资料辨E/B/C，再对应板1E/2B/3C。不能只看外壳平面；需交叉引脚时先套绝缘管并保持间隙，不强挤管脚。',
'焊J1—J6电源端子，进线口朝板外。焊一脚、看贴平及方向，再焊另脚。端子大铜面吸热较多，应确保焊盘和脚同时润湿，不能仅在脚上结锡球。',
'按接口长度裁排针，169针只统计本PCB，不含模块原排针。接口短段逐个对位；先焊一端，校正垂直，再焊其余脚。写清J号与1脚，JP3—JP10不要先全套跳帽。',
'最后装C1=470μF、C2=1000μF、C3=100μF。负极条纹对应GND，1脚为正；实物脚距不同用短绝缘引线并机械固定，不能靠细脚悬挂大电容。',
'单独完成AO3400A转接板。固定一脚后核方向，再焊其余两脚；按AOS脚图1G/2S/3D追踪到转接排针。给线端写G/S/D，不把商品转接板直接当J28同序插座。',
'放大检查每个焊点和引脚，清除助焊剂残留按其说明处理。拍元件面、焊接面照片；对照BOM逐个打勾，再做63.6断电检查。',
])
heading('电气件按值分拣表',3)
groups=defaultdict(list)
for p in parts:
    if p['kind'] in ('resistor','capacitor','electrolytic','npn','logic'): groups[(p['kind'],p['value'])].append(p)
kindname={'resistor':'电阻','capacitor':'瓷片电容','electrolytic':'电解电容','npn':'三极管','logic':'逻辑芯片'}
table(['器件与值','焊位','位号'],[(kindname[k]+' '+v,str(len(arr)),'、'.join(p['ref'] for p in arr)) for (k,v),arr in groups.items()],[4.5,1.4,12],size=9.5,center=(1,))
para('表中的数量是焊位数，不是无条件实装数。37个电阻位含R37备用偏置；R36/R37只能装一个，其他条件上拉由下一节确定。板外传感器适配、MCP和DFPlayer就地去耦、电机抑制电容仍需保留，不能用本表49个电气焊位代替整机BOM。')

heading('63.3  条件电阻和跳帽先记录再安装',2,True)
para('DNP意为该位置暂不装，不是该信号可以没有上拉。先确认模块已有电阻、电源电压和输出类型，再填写实际选择。每个三针选择端子只用一个跳帽：短1—2或2—3，不能两边都短。')
table(['位置','何时安装或设置','记录'],[
('R8/R9 4.7k','I²C已有模块上拉时核并联值；全无上拉时必须装。只拉3V3_EXT。','_____'),
('R10 10k','MCP RESET默认装；模块已有合适同轨上拉可核后不装。','_____'),
('R11/R12 10k','RC522/FLASH片选默认装；同轨已有适宜上拉可核后不装。','_____'),
('R13/R14 10k','只在FLASH相应脚外露且缺同轨上拉时装。','_____'),
('R27/R28 10k','移除编码器原上拉后才装至CORE；保留模块EXT上拉则两只均不装。','_____'),
('R29 4.7k','DHT低压侧须有效上拉，模块/BSS138已有合适上拉则不重复。','_____'),
('R33 10k','编码器SW缺模块上拉、且未验证MCP内部上拉时装。','_____'),
('R34/R35 10k','MODE/BACK默认安装，低有效。','_____'),
('R36/R37 10k','高有效：R36装/R37空；低有效：R36空/R37装。不得同时装或同时空。','_____'),
], [3.2,12.7,2],size=9.5)
heading('供电与信号跳帽',3)
table(['跳帽','默认状态','运行设置'],[
('JP3 蜂鸣信号','开','1—2为PB9直出；2—3为AHCT 5V输出'),
('JP4 蜂鸣电源','开','按模块额定：1—2=3V3_EXT，2—3=5V'),
('JP5 MPU电源','开','按模块实际VIN选择3.3V或5V；I²C仍为3.3V'),
('JP6 HC05电源','开','按裸模块/带稳压底板规格选电源；UART仍为3.3V'),
('JP7 DHT电源','开','3.3V用1—2；5V用2—3并满足JP8条件'),
('JP8 DHT直连','开','仅3.3V安全分支闭合；5V分支必须开，接J23转换'),
('JP9 ADC连接','闭合','正常装帽；核心单独失电前隔离ADC，见供电边界'),
('JP10 MCP复位','开','只临时短接作复位测试，不长期留帽'),
], [3.5,2.2,12.2],size=9.5,center=(1,))
para('JP1与JP2位号刻意空缺，不是漏焊。即使蜂鸣器不用，也保留R36或R37其中一只，使常使能的AHCT第3路输入有确定电平。低有效模块改偏置前必须断电；软件初始化保持不响电平。')

heading('63.4  板外电源和模块怎样连接',2,True)
fig(47,'47_power.png','板外保险开关降压与三支路配电','2S电池经保险开关到外置降压，5V回PCB后分电机舵机逻辑支路；CORE与EXT不并联')
numbered([
'总开关断开，电池暂不插。保护电池正极经板外保险座和总开关到J1.1，负极到J1.2。J2.1/2去降压IN+/IN-；降压应先离板调到5.00V，测好后再把OUT+/OUT-接J3.1/2。电池8.4V绝不能误接J3。',
'J4.1/2接TB6612 VM和动力GND；J5.1/2接SG90正负电源；J6用于逻辑独立分支，DFPlayer电源从这里用合适线径接出。正负均成对走，不能让电机或舵机回流串经核心地线。',
'主电源走20AWG短线，电机/舵机短支路22AWG，长度或压降增加时用20AWG。电池、降压、电机、舵机和喇叭电流不经普通杜邦线。接口信号才用已有10—20cm杜邦线，插好后适度固定。',
'外置3.3V稳压按J7的1=IN+5V、2=公共GND、3=OUT+3V3_EXT接。核心自己的3V3通过J8.3返回参考和规定上拉；不得与EXT输出相连。两路共地不等于两稳压输出可并联。',
'按核心实物丝印逐根接J8/J9。核心仍从教学底板拆出固定；J8/J9不是整排直插插座。PA13/PA14保留SWD，PB2保留BOOT1，PA15/PB3/PB4须在程序中禁JTAG而保留SWD。',
'MCP成品模块接J10五线、J11七线、J12七线，共19根；地址A0/A1/A2在模块侧接地，INTA/INTB不接。GPA7/GPB7不接线且软件设输出0。每次MCU重启重新初始化MCP，先写输出锁存器再配置方向。',
'逐一接OLED、RC522、FLASH、MPU、编码器等，先测电源选择和信号高电平。RC522的SDA丝印在本方案中是SPI片选CS。模块两组输出不能并接；FLASH未使用的WP/HOLD按实际脚型处理。',
'J33六路输入和J31 BUSY无板载5V转换。确认0/3.3V才直连；无5V上拉的开集输出可用10k拉EXT；5V推挽或已有5V上拉则每路外接NPN适配并记录反相。DHT数据线另按双向BSS138分支，不用NPN替代。',
'急停两组独立NC接J27的1—2和3—4，不能用跳帽冒充开关。J27含最高8.4V，只接急停触点。碰撞两NC串联接J26；MODE/BACK与闭盖NO接J25相应成对端子。',
'蜂鸣按JP3/JP4和R36/R37记录接J30；AO转接板按G/S/D接J28，小白LED正负接J29。DFPlayer仅RX经板上1k和BUSY到J31，模块TX不接；喇叭跨模块SPK1/SPK2，二者均不接地。',
])
fig(48,'48_modules.png','核心MCP与已有模块连接总览','核心与MCP和外置模块的J号线束分配，按模块信号名逐线连接')
fig(49,'49_conditions.png','外置条件适配与选择跳帽','5V数字DO的NPN反相适配、DHT的BSS138双向转换，以及蜂鸣控制电平和偏置选择')
para('板上3V3_EXT与3V3_CORE没有直接并联，但ADC和通信线仍可在一侧掉电时注入电流。正常使用两路同启同停，不在外设仍上电时单独拔核心电源。核心USB禁用；独立USB台架须断板电源并拔全部核心线束。USB转串口只接J18 GND/RX/TX，不接任何供电脚。')
para('电源设计预算为5V持续3A、短时峰值至少4A；两电机各自5V堵转电流须不大于1A，SG90瞬态按2A预算，电机与舵机动作分时互锁。逻辑/红外/照明/受控音量语音预算0.5A。模块实际电流与温升还需测量；额定3A文字不能替代带载测试。')

heading('63.5  全部板侧接口针序',2)
para('以下表由A1电气定义生成，覆盖33组外接接口160针及8组跳帽21针。针号是PCB编号，顺序不代表商品模块排针排列。点号表示“接口.引脚”，如J15.3。每接一根，核对线两端名字后打勾。')
sections=[('电源接口',range(1,8)),('核心线束',range(8,10)),('MCP与显示惯性模块',range(10,15)),('RFID存储与串口',range(15,19)),('执行器与温湿度',range(19,24)),('编码器按键和急停',range(24,28)),('照明声音和红外',range(28,34))]
for title,nums in sections:
    heading(title,3)
    refs={f'J{i}' for i in nums}
    rows=[(f"{r['Ref']}.{r['PCB_pin']}",r['Net'],r['External_signal_by_label'],'□') for r in connectors if r['Ref'] in refs]
    table(['板侧针号','网络或电源','外部接到什么信号','检查'],rows,[2.3,5.4,8.9,1.3],size=9,center=(3,))
    notes=[]
    for n in nums:
        p=byref[f'J{n}']; notes.append(f"{p['ref']} {p['value']}：{p['notes'].replace('外接已焊AO3400A转接板','先将AO3400A焊好至转接板后外接')}")
    for note in notes: para(note,9.5)
heading('跳帽与测试端针序',3)
rows=[(f"{r['Ref']}.{r['PCB_pin']}",r['Net'],r['External_signal_by_label'],'□') for r in connectors if r['Ref'].startswith('JP')]
table(['板侧针号','网络或电源','含义','检查'],rows,[2.3,5.4,8.9,1.3],size=9,center=(3,))

heading('63.6  断电检查与首次分阶段上电',2,True)
para('检查框留给实际制作填写。下面是应执行的步骤和预期，不能在未测量时写“已通过”。每增加模块先断电，轮子始终架空到执行器单项测试完成。')
numbered([
'空板焊完但U1不插、所有模块不接，先核位号、阻值、焊桥、端子正负和电解极性。万用表检查VBAT/GND、5V/GND、EXT/GND无持续短路；区分电容充电瞬间响与一直短路。CORE与EXT之间不应有直接导通。',
'确认JP3—JP8开，JP9闭，JP10开；R36装/R37空。外置降压离板先调5.00V，再以限流方式接J3，测J4/J5/J6正端约5V。限流值按当前仅板上负载逐级设置，不带电机舵机。',
'断电接3.3V模块，重上电测J7.3为3.20—3.40V，观察静态电流及异常发热。逐个选择供电跳帽时先空载量中心脚，再接模块；未确认额定电压的跳帽保持开。',
'断电插U1，核缺口与1脚方向。先测14=5V、7=GND；用安全3.3V输入检查三个同相缓冲，PA0/PB15默认低，PB9默认不响电平。第4路11脚保持不接。',
'断电接核心及MCP，先下载最小安全固件：PB5与两路电机PWM为0；MCP先写OLATA/OLATB为0，再设置方向IODIRA=0x40、IODIRB=0x7F，最后初始化功能。GPA7/GPB7均为输出低。',
'I²C先100kHz逐块扫描MCP/OLED/MPU，确认SCL/SDA空闲约3.3V；SPI先低速分别读RC522与FLASH身份，再交替访问。不在通信未稳定时进入积分和移动程序。',
'用受控5V信号验证J21.4经分压到PB1约3.125V；用8.4V电池模拟输入测ADC约1.953V。记录计算值与实测值，检查1%阻值及Vref；核心单独失电前先隔离相关电源与信号，JP9只隔离ADC。',
'不接轮电机，在6.0/7.4/8.4V偏置条件验证双NC。正常释放PB5=1时STBY>2.7V；PB5=0时<0.3V。强制PB5=1并按住急停，STBY仍须<0.3V；PB13应由低变高。逐个拔NC1/NC2回路线并记录。',
'测试释放急停、MCU复位、MCP失联、关键模块断线后的静止状态。急停钳位依赖EXT偏置，不能用正常试验证明EXT失电时也完全安全；异常用总开关切电。STBY低为待机滑行，不会机械锁轮或切舵机电源。',
'架空先单轮低PWM测方向和电流，再双轮；SG90先拆杆空载50Hz/1500μs定位，再按机械章节装连杆。电机行驶与开盖动作互锁，记录5V最低值、温升和复位情况。',
'完成单项后按原步骤逐一跑33项功能：卡片、投放、积分掉电恢复、满桶、循迹、蓝牙、故障锁和导出等。硬件接口齐全不代表固件已经写完；每个功能保存串口日志或实测结果。',
])
table(['验收项','预期或动作','实测记录'],[
('外置降压','空载5.00V；带载项目目标4.85—5.15V','_____V'),
('EXT与CORE','各自约3.3V，不直接并联','_____/_____V'),
('ADC','8.4V输入理论1.953V，按实际Vref校准','_____V'),
('超声分压','5.00V输入理论3.125V','_____V'),
('正常允许','PB5=1且NC正常，STBY>2.7V','_____V'),
('急停保持','PB5强制1且急停按下，STBY<0.3V','_____V'),
('NC2反馈','释放0，按下或断线1','_____'),
('恢复行为','复位/释放/重连后不自动运行','□通过 □待修'),
('电源负载','记录最重正常模式电流与最低5V','_____A / _____V'),
('实际功能','按原33项测试，附日志和故障修复','_____'),
], [3.2,9.5,5.2],size=9.5)
heading('63.7  文件使用与检查范围',2)
para('PCB源工程、制造Gerber、接口表和装配BOM放在仓库hardware/current-low-cost/pcb-a/；本章图源同步保存。订裸PCB使用本次A1制造包，焊接使用本章和A1 BOM，不使用旧V2.1的SMT BOM或坐标文件。')
para('本次原理图6页的ERC报告为0错误、0警告；电气定义检查覆盖293个连接端点和1个未接输出。PCB按本项目设定规则运行DRC，报告为0违规、0未连接项、0原理图一致性差异。检查配置与忽略项保留在报告内；Gerber必须与最终发布检查记录核验的文件一致。这些是设计文件检查结果，不代表实物通电或全部固件已经完成。')
para('更新记录 2026年10月12日：采用A方案，复用直插件和模块；新增图45—49及完整接口表，保留330步骤编号；旧洞洞板和五拆六增改为归档；AHCT第三通道固定PB9，新增蜂鸣及按键偏置位置。')

assert [p.text for p in doc.paragraphs if re.match(r'步骤 \d{3}',p.text)]==steps_before
out=HERE/'staging'/NAME; out.parent.mkdir(exist_ok=True,parents=True); doc.save(out)
audit['output']=str(out); audit['output_sha256']=hashlib.sha256(out.read_bytes()).hexdigest()
audit['steps']=len(steps_before); audit['inline_images']=len(doc.inline_shapes); audit['tables']=len(doc.tables)
(HERE/'build_audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'path':str(out),'images':len(doc.inline_shapes),'tables':len(doc.tables),'steps':len(steps_before),'pending_figures':audit.get('pending_figures',[])},ensure_ascii=False))
