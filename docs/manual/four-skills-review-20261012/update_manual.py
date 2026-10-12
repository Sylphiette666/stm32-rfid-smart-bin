"""Surgical update of the existing illustrated manual; retain all 330 step headings."""
from pathlib import Path
import hashlib,json,re,shutil
from docx import Document
from docx.shared import Pt,Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path(__file__).resolve().parent;ROOT=R.parent;M=R/'manual';M.mkdir(exist_ok=True)
NAME='STM32与RFID智能移动垃圾桶_逐步组装与调试实施手册_图解增强版.docx'
source=ROOT/NAME;before=M/'before.docx'
if not before.exists():shutil.copy2(source,before)
assert hashlib.sha256(before.read_bytes()).hexdigest()=='e372c23d2ecc6cd6eccb2e3e2e680ce35942d956c96c20b01de713b023563885'
d=Document(before);oldsteps=[p.text for p in d.paragraphs if re.match(r'步骤 \d{3}',p.text)];oldtables=[[[c.text for c in row.cells] for row in t.rows] for t in d.tables]
changed=[]
for name in ['45_pcb_top.png','46_pcb_bottom.png']:
 old=hashlib.sha256((R/'baseline/diagrams'/name).read_bytes()).hexdigest();new=(R/'working/diagrams'/name).read_bytes();found=0
 for part in d.part.package.parts:
  if part.partname.startswith('/word/media/') and hashlib.sha256(part.blob).hexdigest()==old:
   part._blob=new;found+=1
 assert found==1,(name,found)
 changed.append({'image':name,'before':old,'after':hashlib.sha256(new).hexdigest()})
def fmt(run,bold=False,size=11):
 run.font.name='宋体';run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体');run.font.size=Pt(size);run.bold=bold
def para(text,bold=False):
 p=d.add_paragraph();p.paragraph_format.space_after=Pt(6);p.paragraph_format.line_spacing=1.15;fmt(p.add_run(text),bold);return p
def heading(text,level=2,newpage=False):
 p=d.add_paragraph(text,f'Heading {level}');p.paragraph_format.page_break_before=newpage;return p
def table(headers,rows,widths):
 t=d.add_table(rows=1,cols=len(headers));t.style='Table Grid';t.autofit=False
 for i,w in enumerate(widths):t.columns[i].width=Cm(w)
 for i,h in enumerate(headers):fmt(t.rows[0].cells[i].paragraphs[0].add_run(h),True,10.5)
 repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
 for row in rows:
  cells=t.add_row().cells
  for i,txt in enumerate(row):
   p=cells[i].paragraphs[0];p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4);fmt(p.add_run(txt),size=10.5)
  prop=t.rows[-1]._tr.get_or_add_trPr();prop.append(OxmlElement('w:cantSplit'))
 for row in t.rows:
  for i,cell in enumerate(row.cells):cell.width=Cm(widths[i])
 props=t._tbl.tblPr
 tw=props.find(qn('w:tblW'));tw.set(qn('w:type'),'dxa');tw.set(qn('w:w'),str(int(sum(widths)/2.54*1440)))
 return t
heading('64  四套技能联合审查与接地加强',1,True)
para('本次审查在原PCB A1上增加55个接地过孔并重新铺铜。裸板制造检查通过，接口、元件采购、模块接法和原330步骤不变。图解46已替换为当前底面图；这些小孔由板厂制作，无需插元件或手工焊线。')
heading('64.1  本次具体完成的工作')
table(['技能套件','实际执行内容'],[
 ('Hubert KiCad-skills','连接真实KiCad 10.0.6，通过IPC加入GND过孔，重新铺铜并保存。'),
 ('mash kicad-skills','使用结构化工具检查原理图、PCB一致性、DNP和规则；没有文本替换CAD。'),
 ('kicad-happy','执行原理图、PCB、EMC、交叉和Gerber分析；核对原厂手册，并实际仿真电池和Echo两处分压。'),
 ('Toaster PCB-design','按布局、回流、电源、工艺及阶段清单检查，记录误报、未满足项与实物验收边界。')
],[4.3,13.4])
para('最终检查：ERC 0违规；PCB DRC 0违规、0未连接、0原理图差异。94个位号逐针一致，33段计划中的粗电源线全部保留。最终共有117个过孔，钻孔文件的411个金属化孔与4个安装孔逐一匹配。')
para('两个分压网络按1%电阻误差做直流仿真：电池8.4V时ADC标称约1.953V、最坏高值约1.984V；Echo在5.15V输入时最坏高值约3.243V。该仿真未包含MCU钳位、线缆寄生或模块动态行为。')
para('本次不增加元件BOM，也不增加焊点。裸板价格仍由实际订单、优惠和运费决定；“多55个过孔”不能保证结算绝对不变，但没有新增需采购的器件。')
heading('64.2  首次调试时追加核对',2,True)
table(['顺序','具体操作与预期'],[
 ('1  认当前文件','下载本次更新的Gerber_RS274X.zip，不能混用旧制造图层。图解46中的小地孔不用焊锡填满。'),
 ('2  仍按原步骤','先执行第63章断电检查和分阶段上电，再接模块；不得直接插上所有模块后带电试错。'),
 ('3  测两路3.3V','J8.3对J8.2测CORE约3.3V；J7.3对J7.2测EXT约3.3V。两路稳压输出不可短接。'),
 ('4  保持线束短','RC522每段信号线建议≤10cm，SCK与地线贴近；MCP的J10地线与J11/J12信号线束一起走，远离电机线。'),
 ('5  先用低速通信','若PCLK2=72MHz，SPI1可先分频256即281.25kHz；I²C先100kHz。低时钟不等于慢边沿，也不能修复地铜缺口。'),
 ('6  验证动态情况','按手册先架空车轮测试急停及锁存，再测试电机启动、舵机动作时是否复位、RFID是否断连。出现问题先测供电和波形。')
],[3.2,14.5])
heading('64.3  仍需实物验证的部分')
para('本版不宣称EMC检查全通过：PA6等信号下方仍缺少连续地铜，部分换层回流需绕行；C8到U1供电脚约5.24mm。55个地过孔有所改善，但未彻底重排信号区。尚不能据此承诺高速SPI、长线或电机强干扰环境稳定。')
para('热分析工具因缺少完整功耗和模块型号，未评估实际温升；生命周期工具也未完成物料查询。这些不能记作已通过。已购端子外形、TO-92针序、电容尺寸仍须装前实测；电机堵转、降压能力、舵机负载和急停动作按原验收步骤验证。')
para('完整报告随设计包保存在 review/four-skills/四套技能联合审查报告.md，包含原始结果、误报依据、原厂资料链接、前后对比和逐项清单。打样属于教学原型验证，尚未完成实物装配、运行或实验室认证。')
assert [p.text for p in d.paragraphs if re.match(r'步骤 \d{3}',p.text)]==oldsteps and len(oldsteps)==330
assert [[[c.text for c in row.cells] for row in t.rows] for t in d.tables[:len(oldtables)]]==oldtables
out=M/NAME;d.save(out)
report={'source_sha256':hashlib.sha256(before.read_bytes()).hexdigest(),'steps':330,'step_headings_preserved':True,'original_tables_preserved':True,'tables':len(d.tables),'images':len(d.inline_shapes),'changed_images':changed,'added_chapter':64,'output':str(out)}
(M/'edit-audit.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8');print(json.dumps(report,ensure_ascii=False))
