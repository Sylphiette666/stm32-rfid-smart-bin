from pathlib import Path
from collections import Counter
import json,hashlib,re,zipfile
from PIL import Image
from docx import Document
R=Path(__file__).resolve().parent;M=R/'manual';N='STM32与RFID智能移动垃圾桶_逐步组装与调试实施手册_图解增强版.docx'
old=Document(M/'before.docx');new=Document(M/N)
steps=lambda d:[p.text for p in d.paragraphs if re.match(r'步骤 \d{3}',p.text)]
tabs=lambda d:[[[c.text for c in row.cells] for row in t.rows] for t in d.tables]
assert steps(old)==steps(new) and len(steps(new))==330
assert tabs(old)==tabs(new)[:len(old.tables)]
def media(p):
 with zipfile.ZipFile(p) as z:return Counter(hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/'))
om=media(M/'before.docx');nm=media(M/N);updates=json.loads((M/'edit-audit.json').read_text())['changed_images']
for u in updates:
 om[u['before']]-=1
 assert nm[u['after']]>=1
assert all(nm[k]>=v for k,v in om.items())
def bodyhash(p):
 im=Image.open(p).convert('RGB');return hashlib.sha256(im.crop((25,70,im.width-25,im.height-70)).tobytes()).hexdigest()
prior=R.parent/'.pcb-lowcost-20261012/a-manual/render_release2';current=M/'render';same=[];changed=[]
for p in sorted(current.glob('page-*.png'),key=lambda p:int(p.stem.split('-')[1])):
 n=int(p.stem.split('-')[1]);(same if (prior/p.name).exists() and bodyhash(p)==bodyhash(prior/p.name) else changed).append(n)
assert changed==[5,179,199,200],changed
assert len(same)==196 and len(new.inline_shapes)==57 and len(new.tables)==66
wr=json.loads((M/'word-render.json').read_text(encoding='utf-8-sig'));assert wr['status']=='success' and wr['pages']==200 and not wr['user_instances_touched']
out={'status':'PASS','docx_sha256':hashlib.sha256((M/N).read_bytes()).hexdigest(),'pdf_sha256':hashlib.sha256((current/'manual.pdf').read_bytes()).hexdigest(),'pages':200,'steps_preserved':330,'existing_table_contents_preserved':64,'total_tables':66,'inline_images':57,'all_other_media_preserved':True,'unchanged_page_bodies':196,'changed_or_added_pages':changed,'visually_inspected_at_full_size':[5,179,199,200],'visual_result':'PASS after correcting added-table column/cell widths; no clipping, overlap or missing text on changed pages.196 other page bodies pixel-identical to reviewed baseline.','render_backend':'Installed Microsoft Word isolated owned instance -> PDF -> pypdfium2 page PNG. Packaged render_docx.py attempted; LibreOffice absent, logged.','user_word_instances_touched':False}
(M/'FINAL_QA.json').write_text(json.dumps(out,ensure_ascii=False,indent=2),encoding='utf-8');print(json.dumps(out,ensure_ascii=False))
